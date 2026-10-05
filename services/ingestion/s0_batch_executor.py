from __future__ import annotations

import argparse
import asyncio
import copy
import csv
import fcntl
import hashlib
import html
import json
import os
import re
import shutil
import subprocess
import tempfile
import time
import unicodedata
from collections.abc import Callable
from contextlib import asynccontextmanager, suppress
from dataclasses import dataclass
from datetime import datetime
from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser
from pathlib import Path, PurePosixPath
from typing import Any

import yaml

from services.ingestion.markdown_evidence import markdown_evidence
from services.ingestion.source_preflight import inspect_source
from services.ingestion.unstructured_router import EXECUTABLE_EXTENSIONS, route_file
from services.ontology_engineering import OntologyWorkflowService
from services.ontology_engineering.storage import MinioArtifactStore
from services.structured_data.pipeline import StructuredDataPipeline, dump_handoff

PAGE_PATTERN = re.compile(rb"/Type\s*/Page(?!s)\b")
PDF_TEXT_EXTRACTOR = Path(__file__).with_name("pdf_text_extract.swift")
PDF_TEXT_BINARY_ROOT = Path("/private/tmp/orion-pdf-text-extractor")
DEFAULT_OCR_RENDER_DPI = 150
DEFAULT_STRUCTURE_RENDER_DPI = 200
MIN_RENDER_DPI = 120
MAX_RENDER_DPI = 300
S0_CACHE_SCHEMA_VERSION = 1
S0_EXECUTOR_REVISION = "2026-09-21-markdown-evidence-spans-v3"
WORKER_EXECUTION_LOCK_TIMEOUT_SECONDS = 60.0
WORKER_EXECUTION_LOCK_POLL_SECONDS = 0.1
ACTIVE_WORKER_RUNNER_STATES = frozenset({"STARTING", "RUNNING"})
ACTIVE_DOCUMENT_JOB_STATUSES = frozenset({"QUEUED", "RUNNING"})
STRUCTURED_FILENAME_ILLEGAL = re.compile(r'[\x00-\x1f<>:"/\\|?*]+')
STRUCTURED_FILENAME_MAX_BYTES = 220
MARKDOWN_TABLE_SEPARATOR = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*$")
MULTI_COLUMN_WHITESPACE = re.compile(r"\S(?:\s{2,}|\t+)\S")
UNSAFE_TEXT_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


@dataclass(frozen=True)
class S0OriginalSnapshot:
    bucket: str
    object_key: str
    source_uri: str
    source_sha256: str
    source_size_bytes: int
    verified_at: str

    def evidence_fields(self) -> dict[str, Any]:
        return {
            "original_storage_backend": "minio",
            "original_bucket": self.bucket,
            "original_object_key": self.object_key,
            "original_source_uri": self.source_uri,
            "original_snapshot_sha256": self.source_sha256,
            "original_snapshot_size_bytes": self.source_size_bytes,
            "original_snapshot_verified_at": self.verified_at,
        }


METHOD_LABELS = {
    "DIRECT_TEXT": "文字层直接解析",
    "HYBRID_PAGE_ROUTE": "逐页混合解析",
    "STRUCTURE_REQUIRED": "需要版面结构识别",
    "PDF_TEXT_LAYER_DIRECT": "PDF 原生文字层直接解析",
    "PADDLEOCR_PP_STRUCTURE_V3": "PaddleOCR 版面结构识别",
    "PADDLEOCR_OCR_FALLBACK": "PaddleOCR 基础文字识别回退",
    "OFFICE_OPEN_XML_WORD": "Word 原生结构解析",
    "OFFICE_OPEN_XML_EXCEL": "Excel 原生结构解析",
    "OFFICE_OR_OPEN_DOCUMENT_DIRECT": "Office 原生结构读取",
    "SPREADSHEET_SCHEMA_AND_VALUES": "表格结构与数据读取",
    "DELIMITED_TABLE_DIRECT": "分隔表格直接解析",
    "TEXT_OR_MARKUP_DIRECT": "文本与标记直接解析",
    "EMAIL_MIME_DIRECT": "邮件正文与头信息解析",
    "PADDLEOCR_OCR": "PaddleOCR 图片文字识别",
    "PADDLEOCR_PP_STRUCTURE_FALLBACK": "PaddleOCR 版面结构回退",
    "DUPLICATE_REUSED": "批次内重复文件复用",
}

SOURCE_LABELS = {
    "PDF": "PDF 文档",
    "DOCX": "Word 文档",
    "XLSX": "Excel 工作簿",
    "XLSM": "Excel 含宏工作簿",
    "CSV": "逗号分隔表格",
    "TSV": "制表符分隔表格",
    "TXT": "纯文本",
    "MD": "Markdown 文档",
    "HTML": "网页文档",
    "HTM": "网页文档",
    "XML": "XML 结构化文本",
    "JSON": "JSON 结构化文本",
    "YAML": "YAML 结构化文本",
    "YML": "YAML 结构化文本",
    "EML": "邮件文件",
    "PNG": "PNG 图片",
    "JPG": "JPG 图片",
    "JPEG": "JPEG 图片",
    "TIF": "TIF 扫描图片",
    "TIFF": "TIFF 扫描图片",
    "BMP": "BMP 图片",
    "WEBP": "WebP 图片",
}


def now() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def atomic_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    temporary.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    temporary.replace(path)


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


class StaleWorkerInvocation(RuntimeError):
    """当前进程已被取消、恢复或后继 invocation 取代。"""


class WorkerExecutionLockTimeout(TimeoutError):
    """同一批次的前一 worker 未能在安全窗口内释放执行锁。"""


@dataclass(frozen=True)
class WorkerInvocationLease:
    request_file: Path
    worker_invocation_id: str
    environment_invocation_id: str

    @classmethod
    def from_request(cls, request_file: Path) -> WorkerInvocationLease:
        request_file = request_file.resolve()
        request = read_json(request_file)
        return cls(
            request_file=request_file,
            worker_invocation_id=str(request.get("worker_invocation_id") or "").strip(),
            environment_invocation_id=str(os.getenv("ORION_WORKER_INVOCATION_ID") or "").strip(),
        )

    @property
    def managed(self) -> bool:
        # 旧 CLI/单元测试没有 invocation 元数据，仍保持原有直接执行语义。
        return bool(self.worker_invocation_id)

    def assert_active(self) -> None:
        if not self.managed:
            return
        try:
            current_request = read_json(self.request_file)
            runner = read_json(self.request_file.parent / "runner.json")
            fence = read_json(self.request_file.parent / "worker-fence.json")
            status = read_json(self.request_file.parent / "status.json")
        except (OSError, ValueError, KeyError) as error:
            raise StaleWorkerInvocation("worker lease 元数据不完整") from error

        expected = self.worker_invocation_id
        observed = {
            self.environment_invocation_id,
            str(current_request.get("worker_invocation_id") or "").strip(),
            str(runner.get("worker_invocation_id") or "").strip(),
            str(fence.get("worker_invocation_id") or "").strip(),
        }
        if observed != {expected}:
            raise StaleWorkerInvocation("worker invocation 已被替换或隔离")
        if str(runner.get("state") or "") not in ACTIVE_WORKER_RUNNER_STATES:
            raise StaleWorkerInvocation("worker runner 已不处于活动状态")
        if str(status.get("status") or "") not in ACTIVE_DOCUMENT_JOB_STATUSES:
            raise StaleWorkerInvocation("批次已经进入终态")


@asynccontextmanager
async def worker_execution_lock(
    job_dir: Path,
    lease: WorkerInvocationLease,
    *,
    timeout_seconds: float = WORKER_EXECUTION_LOCK_TIMEOUT_SECONDS,
    poll_seconds: float = WORKER_EXECUTION_LOCK_POLL_SECONDS,
):
    """用 OS flock 保证同一 job 的旧/新 worker 不会并行写产物。"""

    lock_path = job_dir / ".worker-execution.lock"
    lock_path.parent.mkdir(parents=True, exist_ok=True)
    stream = lock_path.open("a+", encoding="utf-8")
    deadline = time.monotonic() + max(0.0, float(timeout_seconds))
    acquired = False
    try:
        while True:
            lease.assert_active()
            try:
                fcntl.flock(stream.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
                acquired = True
                break
            except BlockingIOError:
                if time.monotonic() >= deadline:
                    raise WorkerExecutionLockTimeout(
                        f"同一批次的旧 worker 未在 {timeout_seconds:g} 秒内释放执行锁。"
                    ) from None
                await asyncio.sleep(max(0.01, float(poll_seconds)))
        lease.assert_active()
        stream.seek(0)
        stream.truncate()
        stream.write(
            json.dumps(
                {
                    "pid": os.getpid(),
                    "worker_invocation_id": lease.worker_invocation_id or None,
                    "acquired_at": now(),
                },
                ensure_ascii=False,
            )
            + "\n"
        )
        stream.flush()
        os.fsync(stream.fileno())
        yield
    finally:
        if acquired:
            fcntl.flock(stream.fileno(), fcntl.LOCK_UN)
        stream.close()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return f"sha256:{digest.hexdigest()}"


def minio_original_snapshot_required() -> bool:
    return os.getenv("ORION_S0_MINIO_SNAPSHOT_REQUIRED", "false").strip().lower() in {
        "1",
        "true",
        "yes",
    }


def prepare_minio_original_snapshots(
    files: list[Path],
    job_dir: Path,
    *,
    store: MinioArtifactStore | None = None,
) -> tuple[
    dict[Path, S0OriginalSnapshot],
    dict[Path, Path],
    list[dict[str, Any]],
]:
    """Snapshot and read back every original before any S0 content inspection."""

    artifact_store = store or MinioArtifactStore.from_env()
    readback_dir = job_dir / ".minio-original-readback"
    readback_dir.mkdir(parents=True, exist_ok=True)
    snapshots: dict[Path, S0OriginalSnapshot] = {}
    processing_paths: dict[Path, Path] = {}
    invocations: list[dict[str, Any]] = []
    for path in files:
        started_at = now()
        source_sha256 = sha256_file(path)
        reference = artifact_store.store_original(path, source_sha256)
        digest = source_sha256.removeprefix("sha256:")
        readback_path = readback_dir / f"{digest}{path.suffix.lower()}"
        artifact_store.client.fget_object(
            artifact_store.bucket,
            reference.object_key,
            str(readback_path),
        )
        observed_sha256 = sha256_file(readback_path)
        if observed_sha256 != source_sha256:
            raise ValueError(f"MinIO 原件回读 SHA-256 不一致：{path.name}；S0 解析已停止。")
        snapshot = S0OriginalSnapshot(
            bucket=artifact_store.bucket,
            object_key=reference.object_key,
            source_uri=f"minio://{artifact_store.bucket}/{reference.object_key}",
            source_sha256=source_sha256,
            source_size_bytes=readback_path.stat().st_size,
            verified_at=now(),
        )
        snapshots[path] = snapshot
        processing_paths[path] = readback_path
        invocations.append(
            {
                "tool": "orion__minio__snapshot_original",
                "tool_version": "ORION S0 MinIO original snapshot 1.0",
                "input": path.name,
                "status": "SUCCEEDED",
                "started_at": started_at,
                "finished_at": snapshot.verified_at,
                **snapshot.evidence_fields(),
            }
        )
    return snapshots, processing_paths, invocations


def filesystem_state(metadata: os.stat_result) -> tuple[int, int, int, int, int]:
    return (
        metadata.st_dev,
        metadata.st_ino,
        metadata.st_size,
        metadata.st_mtime_ns,
        metadata.st_ctime_ns,
    )


def stable_sha256_file(path: Path) -> tuple[str, tuple[int, int, int, int, int]]:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        before = os.fstat(stream.fileno())
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
        after = os.fstat(stream.fileno())
    state = filesystem_state(before)
    if state != filesystem_state(after):
        raise ValueError("REFERENCE 受控资料快照校验失败：文件在哈希期间发生变化")
    return f"sha256:{digest.hexdigest()}", state


def sha256_text(text: str) -> str:
    return f"sha256:{hashlib.sha256(text.encode('utf-8')).hexdigest()}"


def exception_detail(error: BaseException) -> str:
    """展开 ExceptionGroup，避免审计中只留下 TaskGroup 外层错误。"""

    if isinstance(error, BaseExceptionGroup):
        details = [exception_detail(item) for item in error.exceptions]
        return "；".join(dict.fromkeys(detail for detail in details if detail))
    message = str(error).strip()
    return f"{type(error).__name__}: {message}" if message else type(error).__name__


def bounded_dpi(value: Any, default: int) -> int:
    try:
        dpi = int(value)
    except (TypeError, ValueError):
        dpi = default
    return min(MAX_RENDER_DPI, max(MIN_RENDER_DPI, dpi))


def s0_processing_profile(request: dict[str, Any]) -> dict[str, Any]:
    ocr_dpi = bounded_dpi(request.get("ocr_render_dpi"), DEFAULT_OCR_RENDER_DPI)
    structure_dpi = max(
        ocr_dpi,
        bounded_dpi(
            request.get("structure_render_dpi"),
            DEFAULT_STRUCTURE_RENDER_DPI,
        ),
    )
    return {
        "schema_version": S0_CACHE_SCHEMA_VERSION,
        "executor_revision": S0_EXECUTOR_REVISION,
        "ocr_render_dpi": ocr_dpi,
        "structure_render_dpi": structure_dpi,
        "max_excel_rows_per_sheet": int(request.get("max_excel_rows_per_sheet") or 10000),
        "ocr_endpoint": str(request.get("ocr_endpoint") or "http://127.0.0.1:10827/mcp"),
        "structure_endpoint": str(
            request.get("structure_endpoint") or "http://127.0.0.1:10826/mcp"
        ),
    }


def s0_cache_path(root: Path, source_hash: str, request: dict[str, Any]) -> Path:
    source_digest = source_hash.removeprefix("sha256:")
    profile_digest = hashlib.sha256(
        json.dumps(
            s0_processing_profile(request),
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()[:16]
    return (
        root
        / ".orion-s0-cache"
        / f"v{S0_CACHE_SCHEMA_VERSION}"
        / (f"{source_digest}-{profile_digest}.json")
    )


def load_s0_cache(
    root: Path,
    source_hash: str,
    source_suffix: str,
    request: dict[str, Any],
) -> dict[str, Any] | None:
    if request.get("reuse_content_cache") is not True:
        return None
    cache_path = s0_cache_path(root, source_hash, request)
    if not cache_path.is_file():
        return None
    try:
        cached = read_json(cache_path)
    except (OSError, ValueError, json.JSONDecodeError):
        return None
    markdown = str(cached.get("markdown") or "")
    try:
        units = int(cached.get("units"))
    except (TypeError, ValueError):
        return None
    if (
        cached.get("source_sha256") != source_hash
        or cached.get("source_suffix") != source_suffix.lower()
        or cached.get("processing_profile") != s0_processing_profile(request)
        or cached.get("markdown_sha256") != sha256_text(markdown)
        or len(markdown.strip()) < 20
        or not isinstance(cached.get("raw_evidence"), list)
        or not isinstance(cached.get("document_fields"), dict)
        or units <= 0
        or not str(cached.get("method") or "").strip()
    ):
        return None
    return cached


def save_s0_cache(
    root: Path,
    source_hash: str,
    source_suffix: str,
    request: dict[str, Any],
    *,
    markdown: str,
    units: int,
    raw_evidence: list[dict[str, Any]],
    method: str,
    document_fields: dict[str, Any],
    warnings: list[str],
    route_strategy: str,
    route_reason: str,
) -> None:
    if request.get("reuse_content_cache") is not True:
        return
    atomic_json(
        s0_cache_path(root, source_hash, request),
        {
            "schema_version": S0_CACHE_SCHEMA_VERSION,
            "executor_revision": S0_EXECUTOR_REVISION,
            "source_sha256": source_hash,
            "source_suffix": source_suffix.lower(),
            "processing_profile": s0_processing_profile(request),
            "markdown": markdown,
            "markdown_sha256": sha256_text(markdown),
            "units": units,
            "raw_evidence": raw_evidence,
            "method": method,
            "document_fields": document_fields,
            "warnings": warnings,
            "route_strategy": route_strategy,
            "route_reason": route_reason,
            "cached_at": now(),
        },
    )


def structured_markdown_filename(
    source_name: str,
    document_id: str,
    used_names: set[str] | None = None,
) -> str:
    """保留原文件名供人识别，同时用文档编号解决同名冲突。"""

    used_names = used_names if used_names is not None else set()
    source_stem = unicodedata.normalize("NFC", Path(source_name).stem)
    source_stem = STRUCTURED_FILENAME_ILLEGAL.sub("＿", source_stem).strip(" .")
    if not source_stem:
        source_stem = "未命名文档"
    while len(f"{source_stem}.md".encode()) > STRUCTURED_FILENAME_MAX_BYTES:
        source_stem = source_stem[:-1]
    candidate = f"{source_stem}.md"
    normalized = unicodedata.normalize("NFC", candidate).casefold()
    if normalized not in used_names:
        used_names.add(normalized)
        return candidate
    suffix = document_id.removeprefix("DOC-")[-8:] or "DUPLICATE"
    counter = 1
    while True:
        marker = f"__{suffix}" if counter == 1 else f"__{suffix}_{counter}"
        trimmed_stem = source_stem
        while len(f"{trimmed_stem}{marker}.md".encode()) > STRUCTURED_FILENAME_MAX_BYTES:
            trimmed_stem = trimmed_stem[:-1]
        candidate = f"{trimmed_stem}{marker}.md"
        normalized = unicodedata.normalize("NFC", candidate).casefold()
        if normalized not in used_names:
            used_names.add(normalized)
            return candidate
        counter += 1


def rename_job_markdown_outputs(job_dir: Path) -> dict[str, Any]:
    """把旧批次的 DOC-ID 文件迁移为原文件名，并同步所有任务引用。"""

    job_dir = job_dir.resolve()
    bundle_path = job_dir / "bundle.json"
    bundle = read_json(bundle_path)
    documents = bundle.get("documents") or []
    used_names: set[str] = set()
    path_mapping: dict[str, str] = {}
    filename_by_document_id: dict[str, str] = {}
    renamed = 0
    for item in documents:
        document_id = str(item.get("document_id") or "DOC-UNKNOWN")
        old_relative = str(item["structured_markdown_path"])
        new_name = structured_markdown_filename(
            str(item.get("source_name") or document_id),
            document_id,
            used_names,
        )
        old_path = safe_relative(job_dir, old_relative)
        new_path = old_path.with_name(new_name)
        if old_path != new_path:
            if new_path.exists():
                raise FileExistsError(f"目标结构化文件已存在：{new_path.name}")
            old_path.replace(new_path)
            renamed += 1
        new_relative = new_path.relative_to(job_dir).as_posix()
        item["structured_markdown_path"] = new_relative
        item["structured_markdown_filename"] = new_name
        path_mapping[old_relative] = new_relative
        filename_by_document_id[document_id] = new_name

    atomic_json(bundle_path, bundle)
    register_path = job_dir / "document-register.json"
    register = read_json(register_path)
    for item in register:
        document_id = str(item.get("document_id") or "")
        old_relative = str(item.get("structured_markdown_path") or "")
        if old_relative in path_mapping:
            item["structured_markdown_path"] = path_mapping[old_relative]
        if document_id in filename_by_document_id:
            item["structured_markdown_filename"] = filename_by_document_id[document_id]
    atomic_json(register_path, register)

    for status_name in ("status.json", "retry-base-status.json"):
        status_path = job_dir / status_name
        if not status_path.exists():
            continue
        status = read_json(status_path)
        for item in status.get("files") or []:
            old_relative = str(item.get("structured_markdown_path") or "")
            if old_relative in path_mapping:
                item["structured_markdown_path"] = path_mapping[old_relative]
        atomic_json(status_path, status)
    return {
        "job_id": str(bundle.get("processing_trace", {}).get("batch_job_id") or job_dir.name),
        "renamed": renamed,
        "documents": len(documents),
        "directory": str(job_dir / "structured-markdown"),
    }


def safe_relative(root: Path, candidate: str | Path) -> Path:
    root = root.resolve()
    raw = Path(candidate)
    target = (root / raw).resolve() if not raw.is_absolute() else raw.resolve()
    if target != root and root not in target.parents:
        raise ValueError("资料路径超出已批准的接入目录。")
    return target


def reference_snapshot_expectations(
    root: Path,
    source_path: str,
) -> dict[str, Any] | None:
    """Verify an immutable REFERENCE batch before any file is routed or read."""

    root = root.resolve()
    source = safe_relative(root, source_path)
    source_root = source.resolve()
    relative_source = source.relative_to(root)
    if ".orion-s0-uploads" not in relative_source.parts or not source.name.startswith("REFERENCE-"):
        return None

    def fail(reason: str) -> None:
        raise ValueError(f"REFERENCE 受控资料快照校验失败：{reason}")

    raw_source = root / Path(source_path)
    if raw_source.is_symlink() or not source.is_dir():
        fail("快照目录不存在、不是目录或已被替换为符号链接")
    manifest_path = source / "manifest.json"
    if manifest_path.is_symlink() or not manifest_path.is_file():
        fail("manifest.json 不存在、不是文件或已被替换为符号链接")
    try:
        manifest_bytes = manifest_path.read_bytes()
        manifest = json.loads(manifest_bytes.decode("utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as error:
        fail(f"manifest.json 无法读取：{type(error).__name__}")
    if not isinstance(manifest, dict):
        fail("manifest.json 顶层必须是对象")
    if manifest.get("reference_id") != source.name:
        fail("reference_id 与快照目录不一致")
    if manifest.get("source_path") != relative_source.as_posix():
        fail("source_path 与受控接入路径不一致")
    manifest_files = manifest.get("files")
    if not isinstance(manifest_files, list):
        fail("files 必须是数组")

    expectations: dict[Path, dict[str, Any]] = {}
    total_bytes = 0
    for item in manifest_files:
        if not isinstance(item, dict) or set(item) != {"path", "bytes", "sha256"}:
            fail("文件项格式无效")
        raw_path = item.get("path")
        expected_bytes = item.get("bytes")
        expected_sha256 = item.get("sha256")
        if not isinstance(raw_path, str) or not raw_path:
            fail("文件路径为空或格式无效")
        relative = PurePosixPath(raw_path)
        if relative.is_absolute() or ".." in relative.parts or "." in relative.parts:
            fail(f"文件路径不安全：{raw_path}")
        if (
            raw_path == "manifest.json"
            or Path(raw_path).suffix.lower() not in EXECUTABLE_EXTENSIONS
        ):
            fail(f"文件类型不在 S0 正式处理范围：{raw_path}")
        if (
            not isinstance(expected_bytes, int)
            or isinstance(expected_bytes, bool)
            or expected_bytes < 0
        ):
            fail(f"文件大小无效：{raw_path}")
        if not isinstance(expected_sha256, str) or not re.fullmatch(
            r"sha256:[0-9a-f]{64}", expected_sha256
        ):
            fail(f"SHA-256 格式无效：{raw_path}")
        candidate = source.joinpath(*relative.parts)
        cursor = source
        for part in relative.parts:
            cursor /= part
            if cursor.is_symlink():
                fail(f"文件路径包含符号链接：{raw_path}")
        candidate_resolved = candidate.resolve()
        if not candidate.is_file() or (
            candidate_resolved != source_root and source_root not in candidate_resolved.parents
        ):
            fail(f"文件不存在、不是普通文件或越过快照目录：{raw_path}")
        resolved = candidate_resolved
        if resolved in expectations:
            fail(f"文件路径重复：{raw_path}")
        observed_sha256, observed_state = stable_sha256_file(resolved)
        observed_size = observed_state[2]
        if observed_size != expected_bytes or observed_sha256 != expected_sha256:
            fail(f"文件内容与清单不一致：{raw_path}")
        expectations[resolved] = {
            "path": raw_path,
            "bytes": expected_bytes,
            "sha256": expected_sha256,
            "state": observed_state,
        }
        total_bytes += expected_bytes

    actual_files: set[Path] = set()
    for candidate in source.rglob("*"):
        if candidate.is_symlink():
            fail(f"快照中出现符号链接：{candidate.relative_to(source).as_posix()}")
        if candidate.is_file() and candidate.relative_to(source).as_posix() != "manifest.json":
            actual_files.add(candidate.resolve())
    if actual_files != set(expectations):
        fail("实际文件集合与 manifest.json 不一致")
    if manifest.get("file_count") != len(expectations):
        fail("file_count 与实际文件数不一致")
    if manifest.get("total_bytes") != total_bytes:
        fail("total_bytes 与实际文件大小不一致")
    fingerprint = hashlib.sha256()
    for item in sorted(expectations.values(), key=lambda entry: entry["path"]):
        fingerprint.update(
            (
                f"{item['path']}\0{item['bytes']}\0{item['sha256'].removeprefix('sha256:')}\n"
            ).encode()
        )
    expected_reference_id = f"REFERENCE-{fingerprint.hexdigest()[:20].upper()}"
    if source.name != expected_reference_id:
        fail("reference_id 与文件内容指纹不一致")
    # The input root is intentionally mutable (the scheduler creates/removes its
    # lock beside .orion-s0-uploads), so it must not participate in the immutable
    # path-chain state comparison. Still validate it is a real directory.
    if root.is_symlink() or not root.is_dir():
        fail(f"快照根目录不是目录或为符号链接：{root}")
    path_chain: list[dict[str, Any]] = []
    cursor = root
    for relative_part in relative_source.parts:
        cursor /= relative_part
        if cursor.is_symlink() or not cursor.is_dir():
            fail(f"快照路径链包含符号链接或非目录：{cursor}")
        path_chain.append(
            {
                "path": cursor,
                "state": filesystem_state(cursor.stat()),
            }
        )
    return {
        "reference_id": source.name,
        "manifest_sha256": f"sha256:{hashlib.sha256(manifest_bytes).hexdigest()}",
        "root": source,
        "path_chain": path_chain,
        "files": expectations,
    }


def upload_snapshot_expectations(root: Path, source_path: str) -> dict[str, Any] | None:
    """Verify browser upload aliases against their existing private CAS receipts."""
    relative = Path(source_path)
    if (relative.is_absolute() or len(relative.parts) != 2
            or relative.parts[0] != ".orion-s0-uploads"
            or not re.fullmatch(r"UPLOAD-[A-Z0-9-]{8,80}", relative.name)):
        return None
    root = root.resolve()
    source = root / relative
    receipt_root = root / ".orion-s0-content/references/uploads" / relative.name
    directories: dict[Path, dict[str, Any]] = {}

    def fail(reason: str) -> None:
        raise ValueError(f"UPLOAD 受控上传回执校验失败：{reason}")

    def controlled(path: Path, *, directory: bool = False) -> None:
        cursor = root
        for part in path.relative_to(root).parts:
            cursor /= part
            if cursor.is_symlink():
                fail(f"路径包含符号链接：{cursor.relative_to(root)}")
            if cursor != path or directory:
                if not cursor.is_dir():
                    fail(f"目录不存在：{cursor.relative_to(root)}")
                # Shared parents may gain other batches; still pin their identity.
                directories[cursor] = {
                    "path": cursor, "state": filesystem_state(cursor.stat()),
                    "identity_only": not (cursor in (source, receipt_root) or source in cursor.parents),
                }
        if not directory and not path.is_file():
            fail(f"普通文件不存在：{path.relative_to(root)}")

    controlled(source, directory=True)
    controlled(receipt_root, directory=True)
    expectations: dict[Path, dict[str, Any]] = {}
    receipts: list[dict[str, Any]] = []
    verified_cas: set[str] = set()
    for path in sorted(source.rglob("*")):
        if path.is_symlink():
            fail(f"上传批次含符号链接：{path.relative_to(source)}")
        if path.is_dir():
            controlled(path, directory=True)
            continue
        controlled(path)
        name = path.relative_to(source).as_posix()
        if path.suffix.lower() not in EXECUTABLE_EXTENSIONS:
            fail(f"文件类型不在 S0 正式处理范围：{name}")
        receipt_path = receipt_root / (hashlib.sha256(name.encode()).hexdigest() + ".json")
        controlled(receipt_path)
        try:
            receipt_bytes = receipt_path.read_bytes()
            receipt = json.loads(receipt_bytes)
        except (OSError, ValueError) as exc:
            fail(f"上传回执无法读取：{name} ({type(exc).__name__})")
        if not isinstance(receipt, dict) or any((
            receipt.get("schema_version") != 1,
            receipt.get("upload_id") != relative.name,
            receipt.get("batch_source_path") != relative.as_posix(),
            receipt.get("file_name") != name,
        )):
            fail(f"上传回执身份与文件不一致：{name}")
        digest, size = receipt.get("digest"), receipt.get("bytes")
        if (not isinstance(digest, str) or not re.fullmatch(r"sha256:[0-9a-f]{64}", digest)
                or type(size) is not int or size < 0):
            fail(f"上传回执摘要或大小无效：{name}")
        digest_hex = digest.removeprefix("sha256:")
        content_ref = f"sha256/{digest_hex[:2]}/{digest_hex}"
        if receipt.get("content_path") != content_ref:
            fail(f"上传回执 CAS 位置与摘要不一致：{name}")
        content_path = root / ".orion-s0-content" / content_ref
        controlled(content_path)
        if digest not in verified_cas:
            observed_digest, observed_state = stable_sha256_file(content_path)
            if observed_digest != digest or observed_state[2] != size:
                fail(f"CAS 内容与上传回执不一致：{name}")
            verified_cas.add(digest)
        observed_digest, observed_state = stable_sha256_file(path)
        if observed_digest != digest or observed_state[2] != size:
            fail(f"上传文件内容与回执不一致：{name}")
        expectations[path] = {"path": name, "bytes": size, "sha256": digest, "state": observed_state}
        receipts.append({"path": receipt_path, "state": filesystem_state(receipt_path.stat()),
                         "sha256": "sha256:" + hashlib.sha256(receipt_bytes).hexdigest()})
    if not expectations:
        fail("上传批次没有已完成并可验证的文件")
    if set(receipt_root.iterdir()) != {item["path"] for item in receipts}:
        fail("上传文件集合与 CAS 回执集合不一致")
    manifest = [{key: item[key] for key in ("path", "bytes", "sha256")}
                for item in expectations.values()]
    return {
        "source_kind": "UPLOAD", "reference_id": relative.name, "root": source,
        "manifest_sha256": sha256_text(json.dumps(manifest, ensure_ascii=False, sort_keys=True)),
        "path_chain": list(directories.values()),
        "receipt_files": receipts, "files": expectations,
    }


def queued_source_snapshot(root: Path, request: dict[str, Any]) -> dict[str, Any] | None:
    """Recheck the exact submitted manifest before routing a queued batch."""
    frozen = request.get("source_snapshot")
    snapshot = (upload_snapshot_expectations(root, str(request["source_path"]))
                if isinstance(frozen, dict) and frozen.get("kind") == "UPLOAD"
                else reference_snapshot_expectations(root, str(request["source_path"])))
    if frozen is not None and (snapshot is None or snapshot["manifest_sha256"] != frozen.get("manifest_sha256")):
        raise ValueError("受控资料批次与入队时固化的文件清单不一致，不能开始解析。")
    return snapshot


def selected_document_source_scope(
    request: dict[str, Any], *, documents: list[dict[str, Any]] | None = None,
    processing_trace: dict[str, Any] | None = None,
) -> dict[str, Any] | None:
    """Register the exact, platform-verified selection when a queued job creates a project."""
    manifest = request.get("source_snapshot")
    source_prefix = PurePosixPath(str(request["source_path"]))
    if not isinstance(manifest, dict):
        # Historical jobs did not freeze this receipt. Re-read a controlled batch
        # when available; ordinary approved path jobs use their actual processing
        # receipts, restricted to the path the user selected.
        root = Path(request["input_root"]).resolve()
        source_path = str(request["source_path"])
        snapshot = (upload_snapshot_expectations(root, source_path)
                    or reference_snapshot_expectations(root, source_path))
        if snapshot is None:
            if not documents:
                return None
            from services.ontology_engineering.source_scope import (
                normalize_source_scope,
                validate_document_sources,
            )
            selected = safe_relative(root, source_path)
            selection_root = selected.parent if selected.is_file() else selected
            source_prefix = PurePosixPath(selection_root.relative_to(root).as_posix())
            observed = validate_document_sources(
                normalize_source_scope(intake_mode=str(request["intake_mode"])),
                documents, processing_trace=processing_trace,
            )["observed_sources"]
            files = []
            for item in observed:
                path = safe_relative(root, str(item.get("source_path") or ""))
                if ((selected.is_file() and path != selected)
                        or (path != selection_root and selection_root not in path.parents)):
                    raise ValueError("资料执行回执超出用户明确选择的资料路径。")
                files.append({"path": path.relative_to(selection_root).as_posix(),
                              "sha256": item["source_sha256"]})
            manifest = {"files": files}
        else:
            manifest = {"files": list(snapshot["files"].values())}
    sources = []
    for item in manifest.get("files") or []:
        path = PurePosixPath(str(item["path"]))
        if path.is_absolute() or ".." in path.parts or not path.parts:
            raise ValueError("已选择资料清单含无效文件路径。")
        kind = ("STRUCTURED_FILE" if request.get("intake_mode") == "HYBRID"
                and path.suffix.lower() in {".xlsx", ".xlsm", ".xls", ".csv", ".tsv"}
                else "DOCUMENT")
        sources.append({"kind": kind,
                        "source_path": (source_prefix / path).as_posix(),
                        "source_sha256": item["sha256"]})
    if not sources:
        raise ValueError("已选择资料清单为空，不能据此创建工程来源范围。")
    return {"sources": sources}


def assert_reference_snapshot_root(snapshot: dict[str, Any]) -> None:
    root = Path(snapshot["root"])
    try:
        for entry in snapshot["path_chain"]:
            directory = Path(entry["path"])
            if directory.is_symlink() or not directory.is_dir():
                raise ValueError("快照路径链已被替换")
            observed = filesystem_state(directory.stat())
            if ((observed[:2] != entry["state"][:2]) if entry.get("identity_only")
                    else observed != entry["state"]):
                raise ValueError("快照路径链目录身份或状态已变化")
        if snapshot.get("source_kind") == "UPLOAD":
            for receipt in snapshot["receipt_files"]:
                path = Path(receipt["path"])
                if (path.is_symlink() or not path.is_file()
                        or filesystem_state(path.stat()) != receipt["state"]
                        or sha256_file(path) != receipt["sha256"]):
                    raise ValueError("UPLOAD 上传回执身份或内容已变化")
            return
        manifest_path = root / "manifest.json"
        if manifest_path.is_symlink() or not manifest_path.is_file():
            raise ValueError("manifest.json 已被替换")
        observed_manifest_sha256 = sha256_file(manifest_path)
    except (KeyError, OSError, TypeError, ValueError) as error:
        raise ValueError("REFERENCE 受控资料快照校验失败：快照路径链身份或状态已变化") from error
    if observed_manifest_sha256 != snapshot["manifest_sha256"]:
        raise ValueError("REFERENCE 受控资料快照校验失败：manifest.json 内容已变化")


def assert_reference_file(
    path: Path,
    expectations: dict[Path, dict[str, Any]],
    *,
    hash_content: bool = True,
) -> str:
    expected = expectations.get(path.resolve())
    if expected is None:
        raise ValueError("REFERENCE 受控资料快照校验失败：处理文件不在 manifest.json 中")
    if path.is_symlink() or not path.is_file():
        raise ValueError("REFERENCE 受控资料快照校验失败：处理文件已被替换")
    observed = path.stat()
    observed_state = filesystem_state(observed)
    if observed_state != expected["state"]:
        raise ValueError(f"REFERENCE 受控资料快照校验失败：文件内容已变化：{expected['path']}")
    if not hash_content:
        return expected["sha256"]
    observed_sha256, stable_state = stable_sha256_file(path)
    if stable_state != expected["state"] or observed_sha256 != expected["sha256"]:
        raise ValueError(f"REFERENCE 受控资料快照校验失败：文件内容已变化：{expected['path']}")
    return observed_sha256


def collect_files(
    root: Path,
    source_path: str,
    *,
    verified_reference: dict[str, Any] | None = None,
) -> list[Path]:
    source = safe_relative(root, source_path)
    if not source.exists():
        raise ValueError("指定的资料目录或文件不存在。")
    reference_snapshot = verified_reference or reference_snapshot_expectations(root, source_path)
    if reference_snapshot is not None:
        return sorted(reference_snapshot["files"], key=lambda path: path.as_posix())
    candidates = [source] if source.is_file() else sorted(source.rglob("*"))
    source_parts = source.relative_to(root.resolve()).parts
    explicit_upload_batch = ".orion-s0-uploads" in source_parts
    files = [
        path
        for path in candidates
        if path.is_file()
        and not path.is_symlink()
        and path.suffix.lower() in EXECUTABLE_EXTENSIONS
        and ".orion-s0-jobs" not in path.parts
        and ".orion-s0-cache" not in path.parts
        and ".orion-s0-content" not in path.parts
        and (explicit_upload_batch or ".orion-s0-uploads" not in path.parts)
    ]
    if not files:
        raise ValueError("没有找到可处理的 PDF、Word 或 Excel 文件。")
    return files


def load_retry_context(
    root: Path,
    job_dir: Path,
    request: dict[str, Any],
) -> dict[str, Any] | None:
    """Load immutable successes and select only the failed paths for a retry."""
    if request.get("retry_failed_only") is not True:
        return None
    base_status_path = job_dir / "retry-base-status.json"
    if not base_status_path.is_file():
        raise ValueError("缺少失败重试基线，不能安全地跳过成功文件。")
    base_status = read_json(base_status_path)
    failed_by_path = {
        str(item.get("relative_path")): item
        for item in base_status.get("files", [])
        if item.get("status") in {"FAILED", "CANCELLED", "TIMED_OUT"} and item.get("relative_path")
    }
    for relative_path in base_status.get("retryable_paths") or []:
        failed_by_path.setdefault(
            str(relative_path),
            {
                "relative_path": str(relative_path),
                "source_name": Path(str(relative_path)).name,
                "status": base_status.get("status") or "FAILED",
            },
        )
    requested_paths = [str(item) for item in request.get("retry_failed_paths", [])]
    if not requested_paths:
        raise ValueError("没有指定需要重试的失败文件。")
    if len(set(requested_paths)) != len(requested_paths):
        raise ValueError("失败文件重试清单包含重复路径。")
    unexpected = [item for item in requested_paths if item not in failed_by_path]
    if unexpected:
        raise ValueError("失败文件重试清单与上次任务状态不一致。")
    files: list[Path] = []
    for relative_path in requested_paths:
        path = safe_relative(root, relative_path)
        if (
            not path.is_file()
            or path.is_symlink()
            or path.suffix.lower() not in EXECUTABLE_EXTENSIONS
        ):
            raise ValueError(f"失败文件已不存在或不再可处理：{relative_path}")
        files.append(path)
    successful_files = [
        copy.deepcopy(item)
        for item in base_status.get("files", [])
        if item.get("status") == "SUCCEEDED"
    ]
    bundle_path = job_dir / "bundle.json"
    bundle = read_json(bundle_path) if bundle_path.is_file() else {}
    return {
        "base_status": base_status,
        "files": files,
        "failed_by_path": failed_by_path,
        "successful_files": successful_files,
        "bundle": bundle,
    }


def page_count(path: Path) -> int:
    try:
        result = subprocess.run(
            ["pdfinfo", str(path)],
            check=True,
            capture_output=True,
            text=True,
            timeout=20,
        )
        for line in result.stdout.splitlines():
            if line.startswith("Pages:"):
                return max(1, int(line.split(":", 1)[1].strip()))
    except (FileNotFoundError, subprocess.SubprocessError, ValueError):
        pass
    return max(1, len(PAGE_PATTERN.findall(path.read_bytes())))


def recommended_runtime_seconds(total_files: int, total_pages: int) -> int:
    """Return a bounded workload-aware deadline for detached batch execution."""
    estimate = 600 + max(0, total_files) * 5 + max(0, total_pages) * 20
    return min(14400, max(1800, estimate))


def processing_lane(
    path: Path,
    text_layer: dict[str, Any] | None = None,
    *,
    route_category: str | None = None,
) -> str:
    category = route_category or route_file(path).category
    if category == "image":
        return "SPECIALIST"
    if category == "pdf" and not bool((text_layer or {}).get("direct_ready")):
        return "SPECIALIST"
    return "FAST_DIRECT"


def order_processing_lanes(
    files: list[Path],
    pdf_text_layers: dict[Path, dict[str, Any]],
) -> tuple[list[Path], dict[str, int]]:
    fast = [
        path for path in files if processing_lane(path, pdf_text_layers.get(path)) == "FAST_DIRECT"
    ]
    specialist = [
        path for path in files if processing_lane(path, pdf_text_layers.get(path)) == "SPECIALIST"
    ]
    return fast + specialist, {
        "fast_lane_total": len(fast),
        "specialist_lane_total": len(specialist),
        "specialist_concurrency": 1,
    }


def pdf_text_extractor_binary() -> Path:
    source_hash = hashlib.sha256(PDF_TEXT_EXTRACTOR.read_bytes()).hexdigest()[:16]
    PDF_TEXT_BINARY_ROOT.mkdir(parents=True, exist_ok=True)
    binary = PDF_TEXT_BINARY_ROOT / f"pdf-text-extract-{source_hash}"
    if binary.is_file() and os.access(binary, os.X_OK):
        return binary
    environment = {
        **os.environ,
        "SWIFT_MODULECACHE_PATH": "/private/tmp/orion-swift-module-cache",
        "CLANG_MODULE_CACHE_PATH": "/private/tmp/orion-clang-module-cache",
    }
    temporary = binary.with_name(f".{binary.name}.{os.getpid()}.tmp")
    subprocess.run(
        ["swiftc", str(PDF_TEXT_EXTRACTOR), "-o", str(temporary)],
        check=True,
        capture_output=True,
        text=True,
        timeout=90,
        env=environment,
    )
    temporary.chmod(0o700)
    temporary.replace(binary)
    return binary


def extract_pdf_text_layer(path: Path) -> dict[str, Any]:
    environment = {
        **os.environ,
        "SWIFT_MODULECACHE_PATH": "/private/tmp/orion-swift-module-cache",
        "CLANG_MODULE_CACHE_PATH": "/private/tmp/orion-clang-module-cache",
    }
    result = subprocess.run(
        [str(pdf_text_extractor_binary()), str(path)],
        check=True,
        capture_output=True,
        text=True,
        timeout=60,
        env=environment,
    )
    payload = json.loads(result.stdout)
    pages = payload.get("pages") or []
    page_total = int(payload.get("page_count") or len(pages) or page_count(path))
    normalized_pages = []
    for index, item in enumerate(pages, start=1):
        raw_text = str(item.get("text") or "")
        control_character_count = len(UNSAFE_TEXT_CONTROL.findall(raw_text))
        text = UNSAFE_TEXT_CONTROL.sub("", raw_text).strip()
        character_count = int(item.get("character_count") or len(text))
        page_replacements = text.count("�")
        normalized_pages.append(
            {
                "page": int(item.get("page") or index),
                "text": text,
                "character_count": character_count,
                "replacement_ratio": round(page_replacements / max(1, character_count), 6),
                "control_character_count": control_character_count,
                "direct_ready": character_count >= 40
                and page_replacements / max(1, character_count) <= 0.01,
            }
        )
    readable_pages = sum(1 for item in normalized_pages if item["direct_ready"])
    total_characters = sum(int(item.get("character_count") or 0) for item in pages)
    replacement_characters = sum(str(item.get("text") or "").count("�") for item in pages)
    readable_ratio = readable_pages / max(1, page_total)
    replacement_ratio = replacement_characters / max(1, total_characters)
    control_character_count = sum(
        int(item.get("control_character_count") or 0) for item in normalized_pages
    )
    direct_ready = (
        page_total > 0
        and readable_pages == page_total
        and (total_characters >= max(80, page_total * 40) and replacement_ratio <= 0.01)
    )
    markdown = "\n\n".join(
        f"## 第 {int(item.get('page') or index)} 页\n\n{str(item.get('text') or '').strip()}"
        for index, item in enumerate(normalized_pages, start=1)
        if str(item.get("text") or "").strip()
    ).strip()
    return {
        "direct_ready": direct_ready,
        "page_count": page_total,
        "readable_pages": readable_pages,
        "readable_ratio": round(readable_ratio, 4),
        "total_characters": total_characters,
        "replacement_ratio": round(replacement_ratio, 6),
        "control_character_count": control_character_count,
        "markdown": markdown,
        "pages": normalized_pages,
        "reason": (
            f"{readable_pages}/{page_total} 页具有可用文字层，共 {total_characters} 个非空白字符。"
            if direct_ready
            else f"文字层仅覆盖 {readable_pages}/{page_total} 页或质量不足，需要结构化识别。"
        ),
    }


def validate_recognition_image(path: Path) -> dict[str, Any]:
    """确认图片可以被 Pillow 完整解码，避免把截断文件送入 MCP。"""

    if not path.is_file() or path.stat().st_size <= 0:
        raise RuntimeError("识别图片不存在或内容为空。")
    from PIL import Image

    try:
        with Image.open(path) as image:
            image.load()
            width, height = image.size
            image_format = str(image.format or "UNKNOWN")
            mode = str(image.mode)
    except (OSError, SyntaxError, ValueError) as error:
        raise RuntimeError(f"识别图片无法完整解码：{error}") from error
    if width < 32 or height < 32:
        raise RuntimeError(f"识别图片尺寸异常：{width}x{height}。")
    return {
        "format": image_format,
        "mode": mode,
        "width": width,
        "height": height,
        "size_bytes": path.stat().st_size,
    }


def render_pdf_page(
    path: Path,
    page_number: int,
    output_root: Path,
    dpi: int = DEFAULT_OCR_RENDER_DPI,
) -> Path:
    dpi = bounded_dpi(dpi, DEFAULT_OCR_RENDER_DPI)
    last_error: BaseException | None = None
    for attempt in range(1, 3):
        output_prefix = output_root / f"page-{page_number:04d}-{dpi}dpi-{attempt}"
        try:
            subprocess.run(
                [
                    "pdftoppm",
                    "-f",
                    str(page_number),
                    "-l",
                    str(page_number),
                    "-r",
                    str(dpi),
                    "-png",
                    "-singlefile",
                    str(path),
                    str(output_prefix),
                ],
                check=True,
                capture_output=True,
                text=True,
                timeout=120,
            )
            rendered = output_prefix.with_suffix(".png")
            validate_recognition_image(rendered)
            return rendered
        except (OSError, RuntimeError, subprocess.SubprocessError) as error:
            last_error = error
    raise RuntimeError(
        f"PDF 第 {page_number} 页在 {dpi} DPI 下连续两次生成无效识别图片："
        f"{exception_detail(last_error or RuntimeError('未知渲染错误'))}"
    ) from last_error


def normalize_markdown(value: str) -> str:
    text = value.strip()
    try:
        parsed = json.loads(text)
    except (json.JSONDecodeError, TypeError):
        return text

    parts: list[str] = []

    def visit(item: Any) -> None:
        if isinstance(item, str):
            if item.strip():
                parts.append(item.strip())
            return
        if isinstance(item, list):
            for child in item:
                visit(child)
            return
        if isinstance(item, dict):
            preferred = [
                item.get("markdown_texts"),
                item.get("markdown"),
                item.get("text"),
                item.get("content"),
            ]
            if any(candidate for candidate in preferred):
                for candidate in preferred:
                    if candidate:
                        visit(candidate)
                return
            for child in item.values():
                visit(child)

    visit(parsed)
    return "\n\n".join(dict.fromkeys(parts)).strip() or text


def structure_upgrade_reason(markdown: str) -> str | None:
    """Return a narrow reason for upgrading light OCR to layout recognition.

    PP-StructureV3 is intentionally not the default for scanned policy pages: it
    is substantially heavier than text OCR.  Upgrade only when the OCR result
    retains strong table/multi-column signals; an OCR failure is handled by the
    caller as a separate fallback reason.
    """

    lines = [line.rstrip() for line in markdown.splitlines() if line.strip()]
    if any(MARKDOWN_TABLE_SEPARATOR.match(line) for line in lines):
        return "基础 OCR 检测到 Markdown 表格结构"
    pipe_rows = sum(line.count("|") >= 2 for line in lines)
    if pipe_rows >= 3:
        return "基础 OCR 检测到连续表格行"
    multi_column_rows = sum(bool(MULTI_COLUMN_WHITESPACE.search(line)) for line in lines)
    if multi_column_rows >= 4:
        return "基础 OCR 检测到多列对齐内容"
    numeric_rows = sum(
        len(re.findall(r"(?<!\w)\d+(?:[.,%]\d+)?(?!\w)", line)) >= 3 for line in lines
    )
    if numeric_rows >= 4:
        return "基础 OCR 检测到密集数值表格"
    return None


async def call_mcp_tool(
    endpoint: str,
    tool_name: str,
    arguments: dict[str, Any],
) -> tuple[str, dict[str, str]]:
    for proxy_name in (
        "ALL_PROXY",
        "all_proxy",
        "HTTP_PROXY",
        "http_proxy",
        "HTTPS_PROXY",
        "https_proxy",
    ):
        os.environ.pop(proxy_name, None)
    from mcp import ClientSession
    from mcp.client.streamable_http import streamablehttp_client

    async with (
        streamablehttp_client(endpoint) as (read, write, _),
        ClientSession(read, write) as session,
    ):
        initialized = await session.initialize()
        result = await session.call_tool(tool_name, arguments)
        blocks = [str(block.text) for block in result.content if getattr(block, "text", None)]
        text = normalize_markdown("\n\n".join(blocks))
        if getattr(result, "isError", False):
            raise RuntimeError(text or f"{tool_name} 返回失败")
        if len(text.strip()) < 20:
            raise RuntimeError(f"{tool_name} 没有返回足够的结构化文本")
        server_info = getattr(initialized, "serverInfo", None)
        return text, {
            "server_name": str(getattr(server_info, "name", "未提供")),
            "server_version": str(getattr(server_info, "version", "未提供")),
            "protocol_version": str(getattr(initialized, "protocolVersion", "未提供")),
        }


def docx_to_markdown(path: Path) -> tuple[str, int, list[dict[str, Any]]]:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(path)
    parts = [f"# {path.stem}"]
    evidence: list[dict[str, Any]] = []
    unit_count = 0
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            block: Paragraph | Table = Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            block = Table(child, document)
        else:
            continue
        unit_count += 1
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            style = str(block.style.name or "").lower()
            level = (
                1
                if style in {"title", "heading 1"}
                else 2
                if style == "heading 2"
                else 3
                if style == "heading 3"
                else 0
            )
            rendered = f"{'#' * level} {text}" if level else text
            parts.append(rendered)
            evidence.append(
                {
                    "locator": f"Word 段落 {unit_count}",
                    "section": text[:80],
                    "content": text,
                }
            )
            continue
        rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in block.rows]
        if not rows:
            continue
        width = max(len(row) for row in rows)
        rows = [row + [""] * (width - len(row)) for row in rows]
        table_lines = [
            "| " + " | ".join(rows[0]) + " |",
            "|" + "---|" * width,
            *("| " + " | ".join(row) + " |" for row in rows[1:]),
        ]
        parts.extend([f"## 表格 {unit_count}", *table_lines])
        evidence.append(
            {
                "locator": f"Word 表格 {unit_count}",
                "section": f"表格 {unit_count}",
                "content": "\n".join(table_lines),
            }
        )
    markdown = "\n\n".join(part for part in parts if part.strip()).strip()
    return markdown, max(1, unit_count), evidence


def xlsx_to_markdown(
    path: Path,
    max_rows_per_sheet: int,
) -> tuple[str, int, list[dict[str, Any]], list[str]]:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(path, read_only=True, data_only=False)
    parts = [f"# {path.stem}"]
    evidence: list[dict[str, Any]] = []
    warnings: list[str] = []
    try:
        for worksheet in workbook.worksheets:
            parts.append(f"## 工作表：{worksheet.title}")
            rows: list[list[str]] = []
            truncated = False
            for index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                if index > max_rows_per_sheet:
                    truncated = True
                    break
                values = ["" if value is None else str(value).replace("\n", " ") for value in row]
                while values and values[-1] == "":
                    values.pop()
                if values:
                    rows.append(values)
            if truncated:
                warnings.append(
                    f"工作表 {worksheet.title} 超过 {max_rows_per_sheet} 行；当前证据包仅保留前 {max_rows_per_sheet} 行，提交前必须人工确认。"
                )
            if not rows:
                parts.append("（空工作表）")
                evidence.append(
                    {
                        "locator": f"工作表：{worksheet.title}；空工作表",
                        "section": f"工作表：{worksheet.title}",
                        "content": "空工作表",
                    }
                )
                continue
            width = max(len(row) for row in rows)
            padded = [row + [""] * (width - len(row)) for row in rows]
            table_lines = [
                "| " + " | ".join(padded[0]) + " |",
                "|" + "---|" * width,
                *("| " + " | ".join(row) + " |" for row in padded[1:]),
            ]
            parts.extend(table_lines)
            last_column = get_column_letter(max(1, width))
            last_row = len(rows)
            evidence.append(
                {
                    "locator": f"工作表：{worksheet.title}；单元格：A1:{last_column}{last_row}",
                    "section": f"工作表：{worksheet.title}",
                    "content": "\n".join(table_lines),
                }
            )
    finally:
        workbook.close()
    return "\n\n".join(parts).strip(), len(workbook.sheetnames), evidence, warnings


class _HtmlTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip_depth = 0
        self._heading_level = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        if tag in {"script", "style", "noscript"}:
            self._skip_depth += 1
        if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            self._heading_level = int(tag[1])
        if tag in {"p", "div", "section", "article", "tr", "li", "br"}:
            self.parts.append("\n")
        if tag in {"td", "th"}:
            self.parts.append(" | ")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript"} and self._skip_depth:
            self._skip_depth -= 1
        if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            self._heading_level = 0
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._skip_depth or not data.strip():
            return
        prefix = f"{'#' * self._heading_level} " if self._heading_level else ""
        self.parts.append(prefix + " ".join(data.split()))

    def markdown(self) -> str:
        text = "".join(self.parts)
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        return "\n\n".join(lines)


def read_text_safely(path: Path) -> tuple[str, str]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "big5"):
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace"), "utf-8-replacement"


def text_to_markdown(path: Path) -> tuple[str, int, list[dict[str, Any]], list[str]]:
    text, encoding = read_text_safely(path)
    warnings: list[str] = []
    suffix = path.suffix.lower()
    if suffix in {".html", ".htm"}:
        parser = _HtmlTextExtractor()
        parser.feed(text)
        markdown = f"# {path.stem}\n\n{parser.markdown()}"
    elif suffix == ".json":
        payload = json.loads(text)
        markdown = (
            f"# {path.stem}\n\n```json\n{json.dumps(payload, ensure_ascii=False, indent=2)}\n```"
        )
    elif suffix in {".yaml", ".yml"}:
        payload = yaml.safe_load(text)
        markdown = f"# {path.stem}\n\n```yaml\n{yaml.safe_dump(payload, allow_unicode=True, sort_keys=False)}\n```"
    else:
        markdown = text if suffix == ".md" else f"# {path.stem}\n\n{text.strip()}"
    if encoding == "utf-8-replacement":
        warnings.append("源文件编码无法可靠识别，已替换非法字符，提交前必须人工抽查。")
    line_count = max(1, len(markdown.splitlines()))
    evidence = [
        {
            "locator": f"文本第 1-{line_count} 行",
            "section": "全文",
            "content": markdown,
        }
    ]
    if suffix == ".md":
        evidence = markdown_evidence(markdown)
        return markdown, line_count, evidence, warnings
    return markdown.strip(), line_count, evidence, warnings


def delimited_to_markdown(
    path: Path,
    max_rows: int,
) -> tuple[str, int, list[dict[str, Any]], list[str]]:
    text, encoding = read_text_safely(path)
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    rows: list[list[str]] = []
    truncated = False
    for index, row in enumerate(csv.reader(text.splitlines(), delimiter=delimiter), start=1):
        if index > max_rows:
            truncated = True
            break
        rows.append([str(cell).replace("\n", " ") for cell in row])
    if not rows:
        raise ValueError("表格文本没有可读取的数据行。")
    width = max(len(row) for row in rows)
    padded = [row + [""] * (width - len(row)) for row in rows]
    lines = [
        f"# {path.stem}",
        "",
        "| " + " | ".join(padded[0]) + " |",
        "|" + "---|" * width,
        *("| " + " | ".join(row) + " |" for row in padded[1:]),
    ]
    warnings = []
    if truncated:
        warnings.append(
            f"表格超过 {max_rows} 行；当前证据包仅保留前 {max_rows} 行，提交前必须人工确认。"
        )
    if encoding == "utf-8-replacement":
        warnings.append("源文件编码无法可靠识别，已替换非法字符，提交前必须人工抽查。")
    markdown = "\n".join(lines)
    evidence = [
        {
            "locator": f"数据行 1-{len(rows)}；列 1-{width}",
            "section": "表格数据",
            "content": markdown,
        }
    ]
    return markdown, len(rows), evidence, warnings


def eml_to_markdown(path: Path) -> tuple[str, int, list[dict[str, Any]], list[str]]:
    message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
    body = message.get_body(preferencelist=("plain", "html"))
    content = body.get_content() if body else ""
    if body and body.get_content_type() == "text/html":
        parser = _HtmlTextExtractor()
        parser.feed(content)
        content = parser.markdown()
    attachments = [
        part.get_filename() for part in message.iter_attachments() if part.get_filename()
    ]
    metadata = [
        f"- 主题：{message.get('subject', '—')}",
        f"- 发件人：{message.get('from', '—')}",
        f"- 收件人：{message.get('to', '—')}",
        f"- 时间：{message.get('date', '—')}",
        f"- 附件：{'、'.join(attachments) if attachments else '无'}",
    ]
    markdown = (
        f"# {message.get('subject') or path.stem}\n\n"
        + "\n".join(metadata)
        + f"\n\n## 正文\n\n{content.strip()}"
    )
    warnings = ["邮件包含附件；附件需要递归进入同一资料路由。"] if attachments else []
    evidence = [
        {
            "locator": f"邮件 Message-ID：{message.get('message-id', '未提供')}；正文",
            "section": "邮件正文",
            "content": content,
        }
    ]
    return markdown.strip(), 1 + len(attachments), evidence, warnings


def evidence_entries(
    document_id: str,
    source_sha256: str,
    raw_entries: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    entries = raw_entries or [
        {
            "locator": "结构化输出全文",
            "section": "全文",
            "content": source_sha256,
        }
    ]
    result = []
    for index, item in enumerate(entries, start=1):
        content = str(item.get("content") or item.get("section") or source_sha256)
        result.append(
            {
                "evidence_id": f"EVD-{document_id.removeprefix('DOC-')}-{index:03d}",
                "document_id": document_id,
                "source_locator": str(item.get("locator") or "结构化输出全文"),
                "markdown_section": str(item.get("section") or "全文")[:160],
                "content_sha256": sha256_text(content),
            }
        )
    return result


def report_html(status: dict[str, Any], job_dir: Path) -> None:
    rows = []
    for item in status.get("files", []):
        route_label = METHOD_LABELS.get(
            str(item.get("route_strategy", "")),
            "格式专业解析",
        )
        method_label = METHOD_LABELS.get(
            str(item.get("processing_method", "")),
            "尚未完成",
        )
        source_label = SOURCE_LABELS.get(
            str(item.get("source_type", "")),
            "其他资料",
        )
        status_label = {
            "SUCCEEDED": "处理成功",
            "FAILED": "处理失败",
            "RUNNING": "正在处理",
        }.get(str(item.get("status", "")), "尚未开始")
        rows.append(
            "<tr>"
            f"<td>{html.escape(str(item.get('source_name', '—')))}</td>"
            f"<td>{html.escape(source_label)}</td>"
            f"<td><strong>{html.escape(route_label)}</strong><br><small>{html.escape(str(item.get('route_reason', '—')))}</small></td>"
            f"<td>{html.escape(method_label)}</td>"
            f"<td>{html.escape(str(item.get('unit_count', 0)))}</td>"
            f"<td>{html.escape(status_label)}</td>"
            "</tr>"
        )
    warnings = (
        "".join(f"<li>{html.escape(str(item))}</li>" for item in status.get("warnings", []))
        or "<li>没有需要特别确认的截断或降级项。</li>"
    )
    body = f"""<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><title>S0 资料批处理复核报告</title><style>
body{{font-family:-apple-system,BlinkMacSystemFont,'PingFang SC',sans-serif;margin:0;background:#f3f6fb;color:#17233b}}main{{max-width:1080px;margin:40px auto;padding:36px;background:#fff;border-radius:18px;box-shadow:0 18px 50px #1f4b7a18}}h1{{margin:0 0 10px}}p{{line-height:1.8;color:#52627a}}.metrics{{display:grid;grid-template-columns:repeat(4,1fr);gap:12px;margin:24px 0}}.metrics article{{padding:18px;border-radius:12px;background:#f5f8ff}}.metrics span{{display:block;color:#687992;font-size:13px}}.metrics strong{{display:block;font-size:28px;margin-top:7px}}table{{width:100%;border-collapse:collapse}}th,td{{padding:12px;border-bottom:1px solid #e4eaf3;text-align:left}}th{{color:#5f7089;font-size:13px}}section{{margin-top:28px}}li{{margin:8px 0;line-height:1.6}}code{{font-family:ui-monospace,SFMono-Regular,monospace;background:#eef3fb;padding:3px 6px;border-radius:5px}}</style></head><body><main>
<h1>S0 资料批处理复核报告</h1><p>本报告是写入本体工程前的人工复核材料。只有确认文件、结构化文本、证据位置和工具轨迹后，才会正式创建资料工程并通过 S0。</p>
<div class="metrics"><article><span>发现文件</span><strong>{status.get("total_files", 0)}</strong></article><article><span>处理成功</span><strong>{status.get("completed_files", 0)}</strong></article><article><span>处理失败</span><strong>{status.get("failed_files", 0)}</strong></article><article><span>证据条目</span><strong>{status.get("evidence_count", 0)}</strong></article></div>
<section><h2>逐文件处理结果</h2><table><thead><tr><th>原始文件</th><th>类型</th><th>路由决定与原因</th><th>处理方法</th><th>内容单元</th><th>状态</th></tr></thead><tbody>{"".join(rows)}</tbody></table></section>
<section><h2>需要人工确认</h2><ul>{warnings}</ul></section>
<section><h2>审计说明</h2><p>任务编号：<code>{html.escape(str(status.get("job_id", "—")))}</code>。原文件和结构化输出均保存 SHA-256 内容校验码；MCP 工具调用、开始时间、结束时间与降级原因保存在 <code>processing-trace.json</code>。</p></section>
</main></body></html>"""
    (job_dir / "batch-review.html").write_text(body, encoding="utf-8")


async def process_pdf(
    path: Path,
    relative_path: str,
    request: dict[str, Any],
    progress_callback: Callable[[dict[str, Any]], None] | None = None,
    text_layer: dict[str, Any] | None = None,
) -> tuple[str, int, list[dict[str, Any]], str, list[dict[str, Any]]]:
    text_layer_started = now()
    text_layer = text_layer or extract_pdf_text_layer(path)
    pages = int(text_layer["page_count"])
    structure_endpoint = str(request.get("structure_endpoint") or "http://127.0.0.1:10826/mcp")
    ocr_endpoint = str(request.get("ocr_endpoint") or "http://127.0.0.1:10827/mcp")
    processing_profile = s0_processing_profile(request)
    ocr_render_dpi = int(processing_profile["ocr_render_dpi"])
    structure_render_dpi = int(processing_profile["structure_render_dpi"])
    readable_pages = int(text_layer["readable_pages"])
    routing_decision = (
        "DIRECT_TEXT"
        if text_layer["direct_ready"]
        else "HYBRID_PAGE_ROUTE"
        if readable_pages
        else "STRUCTURE_REQUIRED"
    )
    route_reason = (
        text_layer["reason"]
        if text_layer["direct_ready"]
        else f"{readable_pages}/{pages} 页可直接读取，其余页面按页渲染后进行版面识别。"
        if readable_pages
        else text_layer["reason"]
    )
    invocations: list[dict[str, Any]] = [
        {
            "tool": "orion__pdf__direct_text_extract",
            "tool_version": "ORION 文档路由 1.0.0 / macOS PDFKit",
            "input": relative_path,
            "status": "SUCCEEDED",
            "started_at": text_layer_started,
            "finished_at": now(),
            "routing_decision": routing_decision,
            "routing_reason": route_reason,
            "readable_pages": readable_pages,
            "page_count": pages,
            "output_sha256": sha256_text(str(text_layer.get("markdown") or "")),
        }
    ]
    if text_layer["direct_ready"]:
        for item in text_layer["pages"]:
            if progress_callback:
                progress_callback(
                    {
                        "page": int(item["page"]),
                        "status": "SUCCEEDED",
                        "phase": "DIRECT_TEXT",
                        "method": "DIRECT_TEXT",
                        "message": "文字层直接解析完成",
                    }
                )
        markdown = str(text_layer["markdown"])
        raw_evidence = [
            {
                "locator": f"PDF 第 {item['page']} 页；内嵌文字层直接提取",
                "section": f"第 {item['page']} 页",
                "content": item["text"],
            }
            for item in text_layer["pages"]
        ]
        return markdown, pages, raw_evidence, "PDF_TEXT_LAYER_DIRECT", invocations

    markdown_pages: list[str] = []
    raw_evidence: list[dict[str, Any]] = []
    used_structure = False
    used_ocr = False
    with tempfile.TemporaryDirectory(prefix="orion-pdf-pages-", dir="/private/tmp") as temporary:
        rendered_root = Path(temporary)
        for item in text_layer["pages"]:
            current_page = int(item["page"])
            page_text = str(item["text"])
            page_method = "内嵌文字层直接提取"
            if item["direct_ready"]:
                if progress_callback:
                    progress_callback(
                        {
                            "page": current_page,
                            "status": "RUNNING",
                            "phase": "DIRECT_TEXT",
                            "method": "DIRECT_TEXT",
                            "message": "正在读取 PDF 文字层",
                        }
                    )
            else:
                if progress_callback:
                    progress_callback(
                        {
                            "page": current_page,
                            "status": "RUNNING",
                            "phase": "RENDERING",
                            "method": "STRUCTURE_REQUIRED",
                            "message": "正在生成页面图像",
                        }
                    )
                render_started = now()
                rendered = render_pdf_page(
                    path,
                    current_page,
                    rendered_root,
                    ocr_render_dpi,
                )
                rendered_metadata = validate_recognition_image(rendered)
                invocations.append(
                    {
                        "tool": "orion__pdf__render_page_for_recognition",
                        "tool_version": "ORION 文档路由 1.0.0 / Poppler",
                        "input": f"{relative_path}#page={current_page}",
                        "status": "SUCCEEDED",
                        "started_at": render_started,
                        "finished_at": now(),
                        "resolution_dpi": ocr_render_dpi,
                        "image_validation": "PASSED",
                        "image_width": rendered_metadata["width"],
                        "image_height": rendered_metadata["height"],
                        "image_size_bytes": rendered_metadata["size_bytes"],
                    }
                )
                ocr_started = now()
                if progress_callback:
                    progress_callback(
                        {
                            "page": current_page,
                            "status": "RUNNING",
                            "phase": "TEXT_RECOGNITION",
                            "method": "PADDLEOCR_OCR",
                            "message": "正在进行快速文字识别",
                        }
                    )
                ocr_error: Exception | None = None
                try:
                    page_text, provider = await call_mcp_tool(
                        ocr_endpoint,
                        "ocr",
                        {
                            "input_data": str(rendered),
                            "output_mode": "detailed",
                            "return_images": False,
                            "runtime_params": {
                                "use_doc_orientation_classify": True,
                                "use_doc_unwarping": False,
                                "use_textline_orientation": True,
                            },
                        },
                    )
                    used_ocr = True
                    page_method = "PaddleOCR 快速文字识别"
                    invocations.append(
                        {
                            "tool": "mcp__paddleocr_ocr__ocr",
                            "endpoint": ocr_endpoint,
                            "input": f"{relative_path}#page={current_page}",
                            "status": "SUCCEEDED",
                            "started_at": ocr_started,
                            "finished_at": now(),
                            "output_sha256": sha256_text(page_text),
                            **provider,
                        }
                    )
                except Exception as error:
                    ocr_error = error
                    page_text = ""
                    ocr_error_detail = exception_detail(error)
                    invocations.append(
                        {
                            "tool": "mcp__paddleocr_ocr__ocr",
                            "endpoint": ocr_endpoint,
                            "input": f"{relative_path}#page={current_page}",
                            "status": "FAILED",
                            "started_at": ocr_started,
                            "finished_at": now(),
                            "error": ocr_error_detail,
                        }
                    )

                upgrade_reason = (
                    f"基础 OCR 失败：{exception_detail(ocr_error)}"
                    if ocr_error is not None
                    else structure_upgrade_reason(page_text)
                )
                if upgrade_reason:
                    structure_rendered = rendered
                    if structure_render_dpi > ocr_render_dpi:
                        structure_render_started = now()
                        structure_rendered = render_pdf_page(
                            path,
                            current_page,
                            rendered_root,
                            structure_render_dpi,
                        )
                        structure_metadata = validate_recognition_image(structure_rendered)
                        invocations.append(
                            {
                                "tool": "orion__pdf__render_page_for_structure",
                                "tool_version": "ORION 文档路由 1.1.0 / Poppler",
                                "input": f"{relative_path}#page={current_page}",
                                "status": "SUCCEEDED",
                                "started_at": structure_render_started,
                                "finished_at": now(),
                                "resolution_dpi": structure_render_dpi,
                                "upgrade_reason": upgrade_reason,
                                "image_validation": "PASSED",
                                "image_width": structure_metadata["width"],
                                "image_height": structure_metadata["height"],
                                "image_size_bytes": structure_metadata["size_bytes"],
                            }
                        )
                    structure_started = now()
                    if progress_callback:
                        progress_callback(
                            {
                                "page": current_page,
                                "status": "RUNNING",
                                "phase": "STRUCTURE_RECOGNITION",
                                "method": "PADDLEOCR_PP_STRUCTURE_V3",
                                "message": "检测到复杂版面，正在升级结构识别",
                            }
                        )
                    page_text, provider = await call_mcp_tool(
                        structure_endpoint,
                        "pp_structurev3",
                        {
                            "input_data": str(structure_rendered),
                            "output_mode": "simple",
                            "return_images": False,
                            "runtime_params": {
                                "use_doc_orientation_classify": True,
                                "use_doc_unwarping": False,
                                "use_textline_orientation": False,
                                "use_table_recognition": ocr_error is None,
                                "use_formula_recognition": False,
                                "use_chart_recognition": False,
                                "use_seal_recognition": False,
                            },
                        },
                    )
                    used_structure = True
                    page_method = "PaddleOCR 版面结构识别"
                    invocations.append(
                        {
                            "tool": "mcp__paddleocr__pp_structurev3",
                            "endpoint": structure_endpoint,
                            "input": f"{relative_path}#page={current_page}",
                            "status": "SUCCEEDED",
                            "upgrade_reason": upgrade_reason,
                            "table_recognition": ocr_error is None,
                            "resolution_dpi": structure_render_dpi,
                            "started_at": structure_started,
                            "finished_at": now(),
                            "output_sha256": sha256_text(page_text),
                            **provider,
                        }
                    )
            if progress_callback:
                progress_callback(
                    {
                        "page": current_page,
                        "status": "SUCCEEDED",
                        "phase": "PAGE_COMPLETE",
                        "method": (
                            "DIRECT_TEXT"
                            if item["direct_ready"]
                            else "PADDLEOCR_PP_STRUCTURE_V3"
                            if page_method == "PaddleOCR 版面结构识别"
                            else "PADDLEOCR_OCR"
                        ),
                        "message": "本页结构化结果已生成",
                    }
                )
            markdown_pages.append(f"## 第 {current_page} 页\n\n{page_text.strip()}")
            raw_evidence.append(
                {
                    "locator": f"PDF 第 {current_page} 页；{page_method}",
                    "section": f"第 {current_page} 页",
                    "content": page_text,
                }
            )
    markdown = "\n\n".join(markdown_pages).strip()
    if readable_pages or (used_structure and used_ocr):
        method = "HYBRID_PAGE_ROUTE"
    elif used_structure:
        method = "PADDLEOCR_PP_STRUCTURE_V3"
    elif used_ocr:
        method = "PADDLEOCR_OCR"
    else:
        raise RuntimeError("PDF 没有可直接读取或可识别的页面。")
    return markdown, pages, raw_evidence, method, invocations


async def process_image(
    path: Path,
    relative_path: str,
    request: dict[str, Any],
) -> tuple[str, int, list[dict[str, Any]], str, list[dict[str, Any]]]:
    ocr_endpoint = str(request.get("ocr_endpoint") or "http://127.0.0.1:10827/mcp")
    structure_endpoint = str(request.get("structure_endpoint") or "http://127.0.0.1:10826/mcp")
    started_at = now()
    image_metadata = validate_recognition_image(path)
    invocations: list[dict[str, Any]] = [
        {
            "tool": "orion__image__validate_before_recognition",
            "tool_version": "ORION 文档路由 1.1.0 / Pillow",
            "input": relative_path,
            "status": "SUCCEEDED",
            "started_at": started_at,
            "finished_at": now(),
            "image_validation": "PASSED",
            "image_width": image_metadata["width"],
            "image_height": image_metadata["height"],
            "image_size_bytes": image_metadata["size_bytes"],
        }
    ]
    try:
        markdown, provider = await call_mcp_tool(
            ocr_endpoint,
            "ocr",
            {
                "input_data": str(path),
                "output_mode": "detailed",
                "return_images": False,
                "runtime_params": {
                    "use_doc_orientation_classify": True,
                    "use_doc_unwarping": False,
                    "use_textline_orientation": True,
                },
            },
        )
        invocations.append(
            {
                "tool": "mcp__paddleocr_ocr__ocr",
                "endpoint": ocr_endpoint,
                "input": relative_path,
                "status": "SUCCEEDED",
                "started_at": started_at,
                "finished_at": now(),
                "output_sha256": sha256_text(markdown),
                **provider,
            }
        )
        method = "PADDLEOCR_OCR"
        upgrade_reason = structure_upgrade_reason(markdown)
        if upgrade_reason:
            structure_started = now()
            markdown, provider = await call_mcp_tool(
                structure_endpoint,
                "pp_structurev3",
                {
                    "input_data": str(path),
                    "output_mode": "simple",
                    "return_images": False,
                    "runtime_params": {
                        "use_doc_orientation_classify": True,
                        "use_doc_unwarping": False,
                        "use_textline_orientation": False,
                        "use_table_recognition": True,
                        "use_formula_recognition": False,
                        "use_chart_recognition": False,
                        "use_seal_recognition": False,
                    },
                },
            )
            invocations.append(
                {
                    "tool": "mcp__paddleocr__pp_structurev3",
                    "endpoint": structure_endpoint,
                    "input": relative_path,
                    "status": "SUCCEEDED",
                    "upgrade_reason": upgrade_reason,
                    "table_recognition": True,
                    "started_at": structure_started,
                    "finished_at": now(),
                    "output_sha256": sha256_text(markdown),
                    **provider,
                }
            )
            method = "PADDLEOCR_PP_STRUCTURE_V3"
    except Exception as ocr_error:
        ocr_error_detail = exception_detail(ocr_error)
        invocations.append(
            {
                "tool": "mcp__paddleocr_ocr__ocr",
                "endpoint": ocr_endpoint,
                "input": relative_path,
                "status": "FAILED",
                "started_at": started_at,
                "finished_at": now(),
                "error": ocr_error_detail,
            }
        )
        fallback_started = now()
        markdown, provider = await call_mcp_tool(
            structure_endpoint,
            "pp_structurev3",
            {
                "input_data": str(path),
                "output_mode": "simple",
                "return_images": False,
                "runtime_params": {
                    "use_doc_orientation_classify": True,
                    "use_doc_unwarping": False,
                    "use_textline_orientation": False,
                    "use_table_recognition": False,
                    "use_formula_recognition": False,
                    "use_chart_recognition": False,
                    "use_seal_recognition": False,
                },
            },
        )
        invocations.append(
            {
                "tool": "mcp__paddleocr__pp_structurev3",
                "endpoint": structure_endpoint,
                "input": relative_path,
                "status": "SUCCEEDED",
                "fallback_reason": ocr_error_detail,
                "started_at": fallback_started,
                "finished_at": now(),
                "output_sha256": sha256_text(markdown),
                **provider,
            }
        )
        method = "PADDLEOCR_PP_STRUCTURE_FALLBACK"
    evidence = [
        {
            "locator": "图片全文与识别区域",
            "section": "图片识别文本",
            "content": markdown,
        }
    ]
    return markdown, 1, evidence, method, invocations


async def _run_job(request_file: Path, lease: WorkerInvocationLease) -> None:
    request = read_json(request_file)
    actor = request.get("actor")
    job_dir = request_file.parent
    status_path = job_dir / "status.json"
    root = Path(request["input_root"]).resolve()
    reference_snapshot = queued_source_snapshot(root, request)
    reference_expectations = reference_snapshot["files"] if reference_snapshot is not None else {}
    job_id = str(request["job_id"])
    snapshot_readback_dir = job_dir / ".minio-original-readback"
    retry_context = load_retry_context(root, job_dir, request)
    base_status = retry_context["base_status"] if retry_context else {}
    successful_files = retry_context["successful_files"] if retry_context else []
    preserved_bundle = retry_context["bundle"] if retry_context else {}
    status: dict[str, Any] = {
        "job_id": job_id,
        "status": "RUNNING",
        "message": (
            f"已保留 {len(successful_files)} 个成功文件，正在仅重试失败文件。"
            if retry_context
            else "正在发现并处理资料。"
        ),
        "created_at": request.get("created_at") or now(),
        "started_at": now(),
        "source_path": request["source_path"],
        "intake_mode": str(request.get("intake_mode") or "DOCUMENT_ONLY"),
        "retry_failed_only": bool(retry_context),
        "total_files": int(base_status.get("total_files") or 0),
        "completed_files": len(successful_files),
        "failed_files": 0,
        "total_pages": int(base_status.get("total_pages") or 0),
        "processed_pages": sum(int(item.get("processed_pages") or 0) for item in successful_files),
        "current_file": None,
        "current_file_index": 0,
        "current_page": None,
        "current_file_pages": 0,
        "progress_phase": "DISCOVERING",
        "evidence_count": len(preserved_bundle.get("evidence_index") or []),
        # Keep the preserved-success snapshot separate from the mutable status list.
        # Otherwise initial runs append into ``successful_files`` itself and the
        # displayed positions advance as 1, 3, 5... instead of 1, 2, 3....
        "files": copy.deepcopy(successful_files),
        "warnings": copy.deepcopy(base_status.get("warnings") or []),
    }

    def write_status() -> None:
        lease.assert_active()
        atomic_json(status_path, status)

    def write_job_json(path: Path, payload: Any) -> None:
        lease.assert_active()
        atomic_json(path, payload)

    try:
        if (not isinstance(actor, str) or not actor.strip() or len(actor) > 128
                or any(ord(character) < 32 or ord(character) == 127 for character in actor)):
            raise ValueError("资料处理需要请求中明确的操作人标识。")
        actor = actor.strip()
        write_status()
        development_faults_enabled = os.getenv("ORION_ALLOW_DEVELOPMENT_FAULTS") == "1"
        development_delay_ms = (
            min(15000, max(0, int(request.get("development_delay_ms") or 0)))
            if development_faults_enabled
            else 0
        )
        if development_delay_ms:
            status["progress_phase"] = "DEVELOPMENT_DELAY"
            status["message"] = "安全开发态延迟已启用，用于取消/超时验收。"
            write_status()
            await asyncio.sleep(development_delay_ms / 1000)
        files = (
            retry_context["files"]
            if retry_context
            else collect_files(
                root,
                str(request["source_path"]),
                verified_reference=reference_snapshot,
            )
        )
        if reference_snapshot is not None:
            assert_reference_snapshot_root(reference_snapshot)
        source_preflight = inspect_source(
            root / str(request["source_path"]),
            max_document_rows=int(request.get("max_excel_rows_per_sheet") or 10_000),
        )
        status["source_preflight"] = source_preflight
        if source_preflight["requires_structured_import"] and status["intake_mode"] != "HYBRID":
            raise ValueError(
                "数据源预检发现大表，当前 DOCUMENT_ONLY 路线已拒绝；"
                "请新建 HYBRID 工程并选择 IMPORT。"
            )
        if reference_snapshot is not None:
            assert_reference_snapshot_root(reference_snapshot)
            if any(path.resolve() not in reference_expectations for path in files):
                raise ValueError("REFERENCE 受控资料快照校验失败：重试文件不在 manifest.json 中")
            status["reference_snapshot"] = {
                "reference_id": reference_snapshot["reference_id"],
                "manifest_sha256": reference_snapshot["manifest_sha256"],
                "verified": True,
            }
        if not retry_context:
            status["total_files"] = len(files)
        status["discovered_paths"] = [path.relative_to(root).as_posix() for path in files]
        original_snapshots: dict[Path, S0OriginalSnapshot] = {}
        processing_paths = {path: path for path in files}
        snapshot_invocations: list[dict[str, Any]] = []
        if minio_original_snapshot_required():
            status["progress_phase"] = "SNAPSHOTTING_ORIGINALS"
            status["message"] = "正在将原件写入 MinIO 并执行 SHA-256 回读校验。"
            write_status()
            (
                original_snapshots,
                processing_paths,
                snapshot_invocations,
            ) = prepare_minio_original_snapshots(files, job_dir)
            status["original_snapshot"] = {
                "required": True,
                "backend": "minio",
                "verified": True,
                "file_count": len(original_snapshots),
            }
            write_status()
        else:
            status["original_snapshot"] = {
                "required": False,
                "verified": False,
                "reason": "ORION_S0_MINIO_SNAPSHOT_REQUIRED is disabled",
            }
        # Once every original has been written to MinIO and read back with the
        # same SHA-256, parsing uses the job-private readback files. The input
        # root is intentionally mutable (the scheduler creates/removes its lock
        # beside .orion-s0-uploads), so continuing to compare the root directory
        # mtime would reject a valid immutable REFERENCE snapshot.
        reference_source_guard_required = reference_snapshot is not None and not original_snapshots
        original_positions = {path: index for index, path in enumerate(files, start=1)}
        pdf_text_layers: dict[Path, dict[str, Any]] = {}
        for path in files:
            if reference_source_guard_required:
                assert_reference_snapshot_root(reference_snapshot)
                assert_reference_file(
                    path,
                    reference_expectations,
                    hash_content=False,
                )
            route = route_file(path)
            if reference_source_guard_required:
                assert_reference_file(path, reference_expectations, hash_content=False)
                assert_reference_snapshot_root(reference_snapshot)
            if route.category != "pdf":
                continue
            # Keep failures in the specialist lane; process_pdf reports the precise error.
            with suppress(OSError, RuntimeError, ValueError, subprocess.SubprocessError):
                pdf_text_layers[path] = extract_pdf_text_layer(processing_paths[path])
            if reference_source_guard_required:
                assert_reference_file(path, reference_expectations, hash_content=False)
                assert_reference_snapshot_root(reference_snapshot)
        pdf_page_counts: dict[Path, int] = {}
        for path in files:
            if reference_source_guard_required:
                assert_reference_snapshot_root(reference_snapshot)
            route = route_file(path)
            if route.category == "pdf":
                pdf_page_counts[path] = (
                    int(
                        pdf_text_layers[path].get("page_count")
                        or page_count(processing_paths[path])
                    )
                    if path in pdf_text_layers
                    else page_count(processing_paths[path])
                )
            if reference_source_guard_required:
                assert_reference_file(path, reference_expectations, hash_content=False)
                assert_reference_snapshot_root(reference_snapshot)
        if not retry_context:
            status["total_pages"] = sum(pdf_page_counts.values())
        if reference_source_guard_required:
            assert_reference_snapshot_root(reference_snapshot)
        files, queue_summary = order_processing_lanes(files, pdf_text_layers)
        if reference_source_guard_required:
            assert_reference_snapshot_root(reference_snapshot)
        queue_summary["processing_profile"] = s0_processing_profile(request)
        queue_summary["content_cache_enabled"] = request.get("reuse_content_cache") is True
        status["queue_summary"] = queue_summary
        status["recommended_runtime_seconds"] = recommended_runtime_seconds(
            int(status["total_files"]), int(status["total_pages"])
        )
        status["timeout_policy"] = "ADAPTIVE_WORKLOAD"
        status["progress_phase"] = "ROUTING"
        write_status()
        output_dir = job_dir / "structured-markdown"
        output_dir.mkdir(parents=True, exist_ok=True)
        documents: list[dict[str, Any]] = copy.deepcopy(preserved_bundle.get("documents") or [])
        all_evidence: list[dict[str, Any]] = copy.deepcopy(
            preserved_bundle.get("evidence_index") or []
        )
        invocations: list[dict[str, Any]] = copy.deepcopy(
            preserved_bundle.get("processing_trace", {}).get("tool_invocations") or []
        )
        invocations.extend(snapshot_invocations)
        if reference_snapshot is not None:
            invocations.append(
                {
                    "tool": ("orion__upload__verify_cas" if reference_snapshot.get("source_kind") == "UPLOAD"
                             else "orion__reference__verify_manifest"),
                    "tool_version": ("ORION UPLOAD CAS receipts 1.0" if reference_snapshot.get("source_kind") == "UPLOAD"
                                     else "ORION REFERENCE manifest 1.0"),
                    "input": reference_snapshot["reference_id"],
                    "status": "SUCCEEDED",
                    "started_at": status["started_at"],
                    "finished_at": now(),
                    "manifest_sha256": reference_snapshot["manifest_sha256"],
                    "file_count": len(reference_expectations),
                }
            )
        if retry_context:
            invocations.append(
                {
                    "tool": "orion__batch__retry_failed_files",
                    "tool_version": "ORION 文档路由 1.0.0",
                    "input": list(request.get("retry_failed_paths") or []),
                    "status": "RUNNING",
                    "started_at": status["started_at"],
                    "preserved_success_files": len(successful_files),
                }
            )
        total_units = sum(
            int(
                item.get("content_unit_count")
                or item.get("sheet_count")
                or item.get("page_count")
                or 0
            )
            for item in documents
        )
        total_pages = sum(int(item.get("page_count") or 0) for item in documents)
        total_sheets = sum(int(item.get("sheet_count") or 0) for item in documents)
        documents_by_hash: dict[str, dict[str, Any]] = {
            str(item["source_sha256"]): item for item in documents if item.get("source_sha256")
        }
        used_output_names = {
            unicodedata.normalize(
                "NFC", Path(str(item.get("structured_markdown_path") or "")).name
            ).casefold()
            for item in documents
            if item.get("structured_markdown_path")
        }

        def make_pdf_progress_callback(
            current_file_status: dict[str, Any],
            current_path: Path,
        ) -> Callable[[dict[str, Any]], None]:
            def update_pdf_progress(event: dict[str, Any]) -> None:
                page_number = int(event["page"])
                pages = current_file_status.get("pages") or []
                if not 1 <= page_number <= len(pages):
                    return
                page_status = pages[page_number - 1]
                previous_status = page_status.get("status")
                page_status.update(event)
                current_status = str(event.get("status") or "RUNNING")
                if current_status in {"SUCCEEDED", "FAILED"} and previous_status not in {
                    "SUCCEEDED",
                    "FAILED",
                }:
                    current_file_status["processed_pages"] = (
                        int(current_file_status.get("processed_pages") or 0) + 1
                    )
                    status["processed_pages"] = int(status.get("processed_pages") or 0) + 1
                current_file_status["current_page"] = page_number
                status["current_page"] = page_number
                status["current_file_pages"] = int(current_file_status.get("page_count") or 0)
                status["progress_phase"] = str(event.get("phase") or "PROCESSING")
                status["message"] = (
                    f"正在处理 {current_path.name}：第 {page_number}/"
                    f"{current_file_status.get('page_count', 0)} 页 · "
                    f"{event.get('message') or '正在生成结构化结果'}"
                )
                write_status()

            return update_pdf_progress

        for index, path in enumerate(files, start=1):
            relative_path = path.relative_to(root).as_posix()
            previous_file_status = (
                retry_context["failed_by_path"].get(relative_path, {}) if retry_context else {}
            )
            original_position = str(previous_file_status.get("position") or "")
            try:
                display_index = int(original_position.split("/", 1)[0])
            except (TypeError, ValueError):
                display_index = original_positions.get(path, len(successful_files) + index)
            if reference_source_guard_required:
                assert_reference_snapshot_root(reference_snapshot)
            route = route_file(path)
            processing_path = processing_paths[path]
            if reference_source_guard_required:
                assert_reference_file(path, reference_expectations, hash_content=False)
                assert_reference_snapshot_root(reference_snapshot)
            identity_started = now()
            snapshot = original_snapshots.get(path)
            source_hash = (
                snapshot.source_sha256
                if snapshot is not None
                else (
                    assert_reference_file(path, reference_expectations)
                    if reference_snapshot is not None
                    else sha256_file(path)
                )
            )
            if reference_source_guard_required:
                assert_reference_snapshot_root(reference_snapshot)
            source_size_bytes = (
                snapshot.source_size_bytes
                if snapshot is not None
                else (
                    int(reference_expectations[path.resolve()]["bytes"])
                    if reference_snapshot is not None
                    else path.stat().st_size
                )
            )
            invocations.append(
                {
                    "tool": "orion__file__sha256_identity",
                    "tool_version": "ORION 文档路由 1.0.0",
                    "input": relative_path,
                    "status": "SUCCEEDED",
                    "started_at": identity_started,
                    "finished_at": now(),
                    "source_sha256": source_hash,
                    "source_size_bytes": source_size_bytes,
                    "detected_format": route.detected_format,
                }
            )
            document_id = f"DOC-{source_hash.split(':', 1)[1][:16].upper()}"
            file_status = {
                "source_name": path.name,
                "relative_path": relative_path,
                "source_type": path.suffix.removeprefix(".").upper(),
                "route_category": route.category,
                "route_strategy": route.strategy,
                "route_reason": route.reason,
                "processing_lane": processing_lane(
                    path,
                    pdf_text_layers.get(path),
                    route_category=route.category,
                ),
                "status": "RUNNING",
                "position": f"{display_index}/{status['total_files']}",
                **(snapshot.evidence_fields() if snapshot is not None else {}),
            }
            if path in pdf_page_counts:
                file_status.update(
                    {
                        "page_count": pdf_page_counts[path],
                        "processed_pages": 0,
                        "current_page": None,
                        "pages": [
                            {
                                "page": page_number,
                                "status": "PENDING",
                                "phase": "WAITING",
                                "message": "等待处理",
                            }
                            for page_number in range(1, pdf_page_counts[path] + 1)
                        ],
                    }
                )
            status["files"].append(file_status)
            status["current_file"] = path.name
            status["current_file_index"] = display_index
            status["current_page"] = None
            status["current_file_pages"] = int(file_status.get("page_count") or 0)
            status["progress_phase"] = "ROUTING"
            status["message"] = (
                f"正在重试 {path.name}（原批次第 {display_index}/{status['total_files']} 个文件）"
                if retry_context
                else f"正在处理 {path.name}（{index}/{len(files)}）"
            )
            write_status()
            partial_quality = {
                "status": "RUNNING",
                "processed_units": total_units,
                "failed_units": status["failed_files"],
                "review_required": True,
            }
            write_job_json(job_dir / "document-register.json", documents)
            write_job_json(job_dir / "evidence-index.json", all_evidence)
            write_job_json(
                job_dir / "processing-trace.json",
                {
                    "run_id": job_id,
                    "batch_job_id": job_id,
                    "status": "RUNNING",
                    "tool_invocations": invocations,
                    "actor": actor,
                },
            )
            write_job_json(
                job_dir / "bundle.json",
                {
                    "documents": documents,
                    "quality_report": partial_quality,
                    "evidence_index": all_evidence,
                    "processing_trace": {"tool_invocations": invocations},
                },
            )

            update_pdf_progress = make_pdf_progress_callback(file_status, path)

            existing_document = documents_by_hash.get(source_hash)
            if existing_document is not None:
                unit_count = int(
                    existing_document.get("content_unit_count")
                    or existing_document.get("sheet_count")
                    or existing_document.get("page_count")
                    or 0
                )
                original_path = str(existing_document["source_path"])
                file_status.update(
                    {
                        "status": "SUCCEEDED",
                        "route_strategy": "DUPLICATE_REUSED",
                        "route_reason": f"内容校验码与 {original_path} 相同，复用已生成的结构化结果。",
                        "processing_method": "DUPLICATE_REUSED",
                        "unit_count": unit_count,
                        "evidence_count": 0,
                        "structured_markdown_path": existing_document["structured_markdown_path"],
                        "reused_from": original_path,
                    }
                )
                invocations.append(
                    {
                        "tool": "orion__batch__sha256_deduplicate",
                        "tool_version": "ORION 文档路由 1.0.0",
                        "input": relative_path,
                        "status": "SUCCEEDED",
                        "started_at": identity_started,
                        "finished_at": now(),
                        "source_sha256": source_hash,
                        "reused_from": original_path,
                    }
                )
                status["completed_files"] += 1
                if file_status.get("pages"):
                    for item in file_status["pages"]:
                        item.update(
                            {
                                "status": "SUCCEEDED",
                                "phase": "DUPLICATE_REUSED",
                                "method": "DUPLICATE_REUSED",
                                "message": "已复用同内容文件的结构化结果",
                            }
                        )
                    duplicate_pages = int(file_status.get("page_count") or 0)
                    file_status["processed_pages"] = duplicate_pages
                    status["processed_pages"] += duplicate_pages
                if reference_source_guard_required:
                    assert_reference_snapshot_root(reference_snapshot)
                write_status()
                continue

            cached_result = load_s0_cache(root, source_hash, path.suffix, request)
            if cached_result is not None:
                markdown = str(cached_result["markdown"])
                units = int(cached_result["units"])
                raw_evidence = list(cached_result["raw_evidence"])
                method = str(cached_result["method"])
                document_fields = dict(cached_result["document_fields"])
                warnings = [str(item) for item in cached_result.get("warnings") or []]
                output_name = structured_markdown_filename(
                    path.name,
                    document_id,
                    used_output_names,
                )
                output_path = output_dir / output_name
                lease.assert_active()
                output_path.write_text(markdown.rstrip() + "\n", encoding="utf-8")
                document = {
                    "document_id": document_id,
                    "source_name": path.name,
                    "source_path": relative_path,
                    "source_type": path.suffix.removeprefix(".").upper(),
                    "source_sha256": source_hash,
                    "source_size_bytes": source_size_bytes,
                    "detected_format": route.detected_format,
                    "route_category": route.category,
                    "processing_method": method,
                    "structured_markdown_path": output_path.relative_to(job_dir).as_posix(),
                    "structured_markdown_filename": output_name,
                    "structured_markdown_sha256": sha256_text(markdown.rstrip() + "\n"),
                    "cache_reused": True,
                    **(snapshot.evidence_fields() if snapshot is not None else {}),
                    **document_fields,
                }
                documents.append(document)
                documents_by_hash[source_hash] = document
                entries = evidence_entries(document_id, source_hash, raw_evidence)
                all_evidence.extend(entries)
                total_units += units
                total_pages += int(document_fields.get("page_count") or 0)
                total_sheets += int(document_fields.get("sheet_count") or 0)
                status["warnings"].extend(warnings)
                status["completed_files"] += 1
                status["evidence_count"] = len(all_evidence)
                file_status.update(
                    {
                        "status": "SUCCEEDED",
                        "route_strategy": "CONTENT_CACHE_REUSED",
                        "route_reason": "内容校验码和处理策略均未变化，复用跨工程结构化结果。",
                        "processing_method": method,
                        "unit_count": units,
                        "evidence_count": len(entries),
                        "structured_markdown_path": document["structured_markdown_path"],
                        "cache_reused": True,
                    }
                )
                if file_status.get("pages"):
                    cached_pages = int(document_fields.get("page_count") or 0)
                    for page_status in file_status["pages"]:
                        page_status.update(
                            {
                                "status": "SUCCEEDED",
                                "phase": "CONTENT_CACHE_REUSED",
                                "method": method,
                                "message": "已复用同内容、同处理策略的结构化结果",
                            }
                        )
                    file_status["processed_pages"] = cached_pages
                    status["processed_pages"] += cached_pages
                invocations.append(
                    {
                        "tool": "orion__batch__content_cache_reuse",
                        "tool_version": f"ORION S0 cache v{S0_CACHE_SCHEMA_VERSION}",
                        "input": relative_path,
                        "status": "SUCCEEDED",
                        "started_at": identity_started,
                        "finished_at": now(),
                        "source_sha256": source_hash,
                        "processing_profile": s0_processing_profile(request),
                        "cached_at": cached_result.get("cached_at"),
                        "output_sha256": sha256_text(markdown),
                    }
                )
                if reference_source_guard_required:
                    assert_reference_snapshot_root(reference_snapshot)
                write_status()
                continue
            processing_started = now()
            try:
                injected_paths = {
                    str(item) for item in request.get("development_fail_once_paths") or []
                }
                if development_faults_enabled and relative_path in injected_paths:
                    marker = (
                        job_dir
                        / "development-fault-markers"
                        / (hashlib.sha256(relative_path.encode("utf-8")).hexdigest() + ".json")
                    )
                    if not marker.exists():
                        write_job_json(
                            marker,
                            {
                                "relative_path": relative_path,
                                "fault": "FAIL_ONCE",
                                "injected_at": now(),
                            },
                        )
                        raise ValueError("安全开发态失败注入：本文件首次处理按验收规则失败。")
                if reference_source_guard_required:
                    assert_reference_file(path, reference_expectations, hash_content=False)
                    assert_reference_snapshot_root(reference_snapshot)
                warnings: list[str] = []
                if route.category == "pdf":
                    markdown, units, raw_evidence, method, calls = await process_pdf(
                        processing_path,
                        relative_path,
                        request,
                        update_pdf_progress,
                        pdf_text_layers.get(path),
                    )
                    document_fields = {"page_count": units}
                    total_pages += units
                    invocations.extend(calls)
                    file_status["route_strategy"] = calls[0]["routing_decision"]
                    file_status["route_reason"] = calls[0]["routing_reason"]
                elif route.category == "word" and route.detected_format == ".docx":
                    markdown, units, raw_evidence = docx_to_markdown(processing_path)
                    method = "OFFICE_OPEN_XML_WORD"
                    document_fields = {"content_unit_count": units}
                    invocations.append(
                        {
                            "tool": "orion__office__docx_to_markdown",
                            "tool_version": "ORION 文档路由 1.0.0 / python-docx",
                            "input": relative_path,
                            "status": "SUCCEEDED",
                            "started_at": processing_started,
                            "finished_at": now(),
                            "output_sha256": sha256_text(markdown),
                            "routing_reason": route.reason,
                        }
                    )
                elif route.category == "spreadsheet" and route.detected_format == ".xlsx":
                    markdown, units, raw_evidence, warnings = xlsx_to_markdown(
                        processing_path,
                        int(request.get("max_excel_rows_per_sheet") or 10000),
                    )
                    method = "OFFICE_OPEN_XML_EXCEL"
                    document_fields = {"sheet_count": units}
                    total_sheets += units
                    invocations.append(
                        {
                            "tool": "orion__office__xlsx_to_markdown",
                            "tool_version": "ORION 文档路由 1.0.0 / openpyxl",
                            "input": relative_path,
                            "status": "SUCCEEDED",
                            "started_at": processing_started,
                            "finished_at": now(),
                            "output_sha256": sha256_text(markdown),
                            "routing_reason": route.reason,
                        }
                    )
                elif route.category == "spreadsheet" and path.suffix.lower() in {".csv", ".tsv"}:
                    markdown, units, raw_evidence, warnings = delimited_to_markdown(
                        processing_path,
                        int(request.get("max_excel_rows_per_sheet") or 10000),
                    )
                    method = "DELIMITED_TABLE_DIRECT"
                    document_fields = {"content_unit_count": units}
                    invocations.append(
                        {
                            "tool": "orion__table__delimited_to_markdown",
                            "tool_version": "ORION 文档路由 1.0.0 / Python csv",
                            "input": relative_path,
                            "status": "SUCCEEDED",
                            "started_at": processing_started,
                            "finished_at": now(),
                            "output_sha256": sha256_text(markdown),
                            "routing_reason": route.reason,
                        }
                    )
                elif route.category == "text" and path.suffix.lower() in {
                    ".txt",
                    ".md",
                    ".html",
                    ".htm",
                    ".xml",
                    ".json",
                    ".yaml",
                    ".yml",
                }:
                    markdown, units, raw_evidence, warnings = text_to_markdown(processing_path)
                    method = "TEXT_OR_MARKUP_DIRECT"
                    document_fields = {"content_unit_count": units}
                    invocations.append(
                        {
                            "tool": "orion__text__direct_to_markdown",
                            "tool_version": "ORION 文档路由 1.0.0",
                            "input": relative_path,
                            "status": "SUCCEEDED",
                            "started_at": processing_started,
                            "finished_at": now(),
                            "output_sha256": sha256_text(markdown),
                            "routing_reason": route.reason,
                        }
                    )
                elif route.category == "email" and path.suffix.lower() == ".eml":
                    markdown, units, raw_evidence, warnings = eml_to_markdown(processing_path)
                    method = "EMAIL_MIME_DIRECT"
                    document_fields = {"content_unit_count": units}
                    invocations.append(
                        {
                            "tool": "orion__email__eml_to_markdown",
                            "tool_version": "ORION 文档路由 1.0.0 / Python email",
                            "input": relative_path,
                            "status": "SUCCEEDED",
                            "started_at": processing_started,
                            "finished_at": now(),
                            "output_sha256": sha256_text(markdown),
                            "routing_reason": route.reason,
                        }
                    )
                elif route.category == "image":
                    markdown, units, raw_evidence, method, calls = await process_image(
                        processing_path, relative_path, request
                    )
                    document_fields = {"content_unit_count": units}
                    invocations.extend(calls)
                else:
                    raise ValueError(f"当前执行器尚未启用 {route.label} 的正式适配器。")
                if len(markdown.strip()) < 20:
                    raise ValueError("结构化 Markdown 内容不足，不能进入人工复核。")
                if reference_source_guard_required:
                    assert_reference_file(
                        path,
                        reference_expectations,
                        hash_content=False,
                    )
                    assert_reference_snapshot_root(reference_snapshot)
                total_units += units
                output_name = structured_markdown_filename(
                    path.name,
                    document_id,
                    used_output_names,
                )
                output_path = output_dir / output_name
                lease.assert_active()
                output_path.write_text(markdown.rstrip() + "\n", encoding="utf-8")
                document = {
                    "document_id": document_id,
                    "source_name": path.name,
                    "source_path": relative_path,
                    "source_type": path.suffix.removeprefix(".").upper(),
                    "source_sha256": source_hash,
                    "source_size_bytes": source_size_bytes,
                    "detected_format": route.detected_format,
                    "route_category": route.category,
                    "processing_method": method,
                    "structured_markdown_path": output_path.relative_to(job_dir).as_posix(),
                    "structured_markdown_filename": output_name,
                    "structured_markdown_sha256": sha256_text(markdown.rstrip() + "\n"),
                    **(snapshot.evidence_fields() if snapshot is not None else {}),
                    **document_fields,
                }
                documents.append(document)
                documents_by_hash[source_hash] = document
                entries = evidence_entries(document_id, source_hash, raw_evidence)
                all_evidence.extend(entries)
                try:
                    lease.assert_active()
                    save_s0_cache(
                        root,
                        source_hash,
                        path.suffix,
                        request,
                        markdown=markdown,
                        units=units,
                        raw_evidence=raw_evidence,
                        method=method,
                        document_fields=document_fields,
                        warnings=warnings,
                        route_strategy=str(file_status.get("route_strategy") or route.strategy),
                        route_reason=str(file_status.get("route_reason") or route.reason),
                    )
                except OSError as cache_error:
                    invocations.append(
                        {
                            "tool": "orion__batch__content_cache_store",
                            "tool_version": f"ORION S0 cache v{S0_CACHE_SCHEMA_VERSION}",
                            "input": relative_path,
                            "status": "FAILED_NON_BLOCKING",
                            "started_at": processing_started,
                            "finished_at": now(),
                            "error": exception_detail(cache_error),
                        }
                    )
                status["warnings"].extend(warnings)
                file_status.update(
                    {
                        "status": "SUCCEEDED",
                        "processing_method": method,
                        "unit_count": units,
                        "evidence_count": len(entries),
                        "structured_markdown_path": document["structured_markdown_path"],
                    }
                )
                status["completed_files"] += 1
                status["evidence_count"] = len(all_evidence)
            except StaleWorkerInvocation:
                raise
            except Exception as error:
                error_detail = exception_detail(error)
                if file_status.get("pages") and file_status.get("current_page"):
                    update_pdf_progress(
                        {
                            "page": int(file_status["current_page"]),
                            "status": "FAILED",
                            "phase": "PAGE_FAILED",
                            "message": error_detail,
                        }
                    )
                file_status.update({"status": "FAILED", "error": error_detail})
                invocations.append(
                    {
                        "tool": "orion__router__process_file",
                        "tool_version": "ORION 文档路由 1.0.0",
                        "input": relative_path,
                        "status": "FAILED",
                        "started_at": processing_started,
                        "finished_at": now(),
                        "route_category": route.category,
                        "error": error_detail,
                    }
                )
                status["failed_files"] += 1
            write_status()

        if reference_source_guard_required:
            assert_reference_snapshot_root(reference_snapshot)

        def file_position(item: dict[str, Any]) -> int:
            try:
                return int(str(item.get("position") or "0").split("/", 1)[0])
            except ValueError:
                return 0

        status["files"].sort(key=file_position)
        if retry_context:
            retry_invocation = next(
                (
                    item
                    for item in reversed(invocations)
                    if item.get("tool") == "orion__batch__retry_failed_files"
                    and item.get("status") == "RUNNING"
                ),
                None,
            )
            if retry_invocation is not None:
                retry_invocation.update(
                    {
                        "status": ("SUCCEEDED" if status["failed_files"] == 0 else "FAILED"),
                        "finished_at": now(),
                        "retried_files": len(files),
                        "remaining_failed_files": status["failed_files"],
                    }
                )

        quality_report = {
            "status": "READY_FOR_REVIEW" if not status["failed_files"] else "FAILED",
            "processed_units": total_units,
            "failed_units": status["failed_files"],
            "low_confidence_units": 0,
            "reviewed_low_confidence_units": 0,
            "unreviewed_low_confidence_units": 0,
            "page_count": total_pages,
            "sheet_count": total_sheets,
            "warning_count": len(status["warnings"]),
            "confidence_scores_available": False,
            "review_required": True,
            "content_cache_reused_documents": sum(
                1 for document in documents if document.get("cache_reused") is True
            ),
        }
        processing_trace = {
            "run_id": job_id,
            "batch_job_id": job_id,
            "router_version": "1.1.0",
            "executor_version": "1.1.0",
            "executor_revision": S0_EXECUTOR_REVISION,
            "provider": "PaddleOCR 双 MCP + Office Open XML 解析器",
            "tool_calls": list(dict.fromkeys(call["tool"] for call in invocations)),
            "tool_invocations": invocations,
            "processing_profile": s0_processing_profile(request),
            "content_cache_enabled": request.get("reuse_content_cache") is True,
            "content_cache_reused_documents": sum(
                1 for document in documents if document.get("cache_reused") is True
            ),
            "actor": actor,
            "source_path": str(request["source_path"]),
            "started_at": status["started_at"],
            "finished_at": now(),
        }
        bundle = {
            "documents": documents,
            "quality_report": quality_report,
            "evidence_index": all_evidence,
            "processing_trace": processing_trace,
            "source_preflight": status.get("source_preflight"),
        }
        write_job_json(job_dir / "document-register.json", documents)
        write_job_json(job_dir / "evidence-index.json", all_evidence)
        write_job_json(job_dir / "ingestion-quality-report.json", quality_report)
        write_job_json(job_dir / "processing-trace.json", processing_trace)
        write_job_json(job_dir / "bundle.json", bundle)
        status["finished_at"] = now()
        status["current_file"] = None
        status["current_page"] = None
        status["current_file_pages"] = 0
        status["progress_phase"] = "COMPLETE"
        status["review_report"] = "batch-review.html"
        status["status"] = "READY_FOR_REVIEW" if status["failed_files"] == 0 else "FAILED"
        status["message"] = (
            "解析完成，等待人工复核后写入 S0。"
            if status["status"] == "READY_FOR_REVIEW"
            else "部分文件处理失败；请查看逐文件原因后重试。"
        )
        lease.assert_active()
        report_html(status, job_dir)
        write_status()
    except StaleWorkerInvocation:
        return
    except Exception as error:
        status.update(
            {
                "status": "FAILED",
                "message": "资料批处理未完成。",
                "error": exception_detail(error),
                "finished_at": now(),
            }
        )
        try:
            lease.assert_active()
            report_html(status, job_dir)
            write_status()
        except StaleWorkerInvocation:
            return
    finally:
        shutil.rmtree(snapshot_readback_dir, ignore_errors=True)


async def run_job(request_file: Path) -> None:
    request_file = request_file.resolve()
    lease = WorkerInvocationLease.from_request(request_file)
    try:
        lease.assert_active()
        async with worker_execution_lock(request_file.parent, lease):
            lease.assert_active()
            await _run_job(request_file, lease)
    except StaleWorkerInvocation:
        # 取消、超时或后继恢复已经取得所有权时，旧进程必须静默退出。
        return


def commit_job(
    job_dir: Path,
    reviewed_by: str,
    rationale: str,
    accept_warnings: bool,
    structured_data_action: str = "DOCUMENT_ONLY",
    *,
    resume_import: bool = False,
) -> dict[str, Any]:
    """Only one commit/import may own a job, including after client timeout."""
    with (job_dir / ".structured-import.lock").open("a+") as lock:
        try:
            fcntl.flock(lock.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError as exc:
            raise ValueError("该资料任务正在提交或导入，请回读状态，不要重复执行。") from exc
        try:
            return _commit_job_unlocked(job_dir, reviewed_by, rationale, accept_warnings,
                                        structured_data_action, resume_import=resume_import)
        finally:
            fcntl.flock(lock.fileno(), fcntl.LOCK_UN)


def _commit_job_unlocked(
    job_dir: Path,
    reviewed_by: str,
    rationale: str,
    accept_warnings: bool,
    structured_data_action: str,
    *,
    resume_import: bool,
) -> dict[str, Any]:
    status_path = job_dir / "status.json"
    status = read_json(status_path)
    if resume_import:
        if (status.get("status") != "COMMITTED" or structured_data_action != "IMPORT"
                or (status.get("structured_data") or {}).get("status") not in {
                    "IMPORTING", "IMPORT_FAILED", "WORKFLOW_GATE_FAILED"}):
            raise ValueError("当前任务不属于可恢复的 S1 结构化导入。")
    elif status.get("status") != "READY_FOR_REVIEW":
        raise ValueError("当前任务尚未完成解析，不能写入 S0。")
    if not rationale.strip():
        raise ValueError("请填写复核结论。")
    if status.get("warnings") and not accept_warnings:
        raise ValueError("当前批次有截断或降级提醒，必须明确确认后才能提交。")
    request = read_json(job_dir / "request.json")
    intake_mode = str(request.get("intake_mode") or "DOCUMENT_ONLY")
    source_preflight = status.get("source_preflight") or {}
    normalized_structured_action = structured_data_action.strip().upper()
    if normalized_structured_action not in {"DOCUMENT_ONLY", "IMPORT"}:
        raise ValueError("structured_data_action 必须是 DOCUMENT_ONLY 或 IMPORT。")
    if normalized_structured_action == "IMPORT" and intake_mode != "HYBRID":
        raise ValueError("只有 HYBRID 工程可以把文件表格导入结构化数据区并进入 S1。")
    if (
        source_preflight.get("requires_structured_import") is True
        and normalized_structured_action != "IMPORT"
    ):
        raise ValueError("数据源预检确认存在大表，必须选择 IMPORT 后才能提交 S0。")
    bundle = read_json(job_dir / "bundle.json")
    documents = []
    for item in bundle["documents"]:
        markdown_path = safe_relative(job_dir, item["structured_markdown_path"])
        documents.append(
            {
                **{key: value for key, value in item.items() if key != "structured_markdown_path"},
                "structured_markdown": markdown_path.read_text(encoding="utf-8"),
            }
        )
    quality_report = {
        **bundle["quality_report"],
        "status": "PASSED",
        "reviewed_by": reviewed_by,
        "review_rationale": rationale.strip(),
        "reviewed_at": now(),
        "warnings_accepted": bool(status.get("warnings") and accept_warnings),
        "accepted_warnings": list(status.get("warnings") or []) if accept_warnings else [],
        "structured_data_action": normalized_structured_action,
    }
    processing_trace = {
        **bundle["processing_trace"],
        "actor": reviewed_by,
        "review_rationale": rationale.strip(),
    }
    workflow_home = Path(request["workflow_home"]).resolve()
    service = OntologyWorkflowService(workflow_home)
    project_id = str(request.get("project_id") or "").strip()
    if not project_id:
        if intake_mode not in {"DOCUMENT_ONLY", "HYBRID"}:
            raise ValueError("资料批处理只支持文件资料建模或混合建模工程。")
        selected_scope = selected_document_source_scope(
            request, documents=documents, processing_trace=processing_trace,
        )
        created = service.create_project(
            project_name=str(request["project_name"]),
            domain=str(request["domain"]),
            intake_mode=intake_mode,
            intake_rationale=str(request["intake_rationale"]),
            cq_mode=str(request.get("cq_mode") or "USER_PLUS_AI"),
            initial_competency_questions=request.get("initial_competency_questions") or [],
            source_scope=selected_scope,
            request_id=str(request.get("project_request_id") or f"document-job:{status['job_id']}"),
        )
        project_id = str(created["project_id"])
    if resume_import:
        project_dir, state = service._require_stage(project_id, "S1")
        service._require_expected_revision(state, (status.get("workflow") or {}).get("revision"))
        if service._verify_project_integrity(project_dir, state, stages=("S0",))["status"] != "PASSED":
            raise ValueError("S0 正式资产未通过完整性回读，不能恢复 S1 导入。")
        result = service.get_status(project_id=project_id)
    else:
        result = service.record_document_evidence(
            project_id=project_id,
            documents=documents,
            quality_report=quality_report,
            evidence_index=bundle["evidence_index"],
            processing_trace=processing_trace,
            expected_revision=request.get("project_revision"),
        )
    structured_data: dict[str, Any] = {
        "action": normalized_structured_action,
        "status": "NOT_REQUESTED",
    }
    if normalized_structured_action == "IMPORT":
        data_import_succeeded = False
        completed_workflow_stages = ["S0"]
        try:
            structured_data["status"] = "IMPORTING"
            status.update(
                {
                    "status": "COMMITTED",
                    "project_id": project_id,
                    "workflow": result,
                    "structured_data": structured_data,
                    "reviewed_by": reviewed_by,
                    "review_rationale": rationale.strip(),
                    "message": "S0 已提交，正在把文件表格全量导入独立结构化数据区。",
                }
            )
            atomic_json(status_path, status)
            handoff = StructuredDataPipeline().import_documents(
                project_id=project_id,
                documents=documents,
                input_root=Path(request["input_root"]).resolve(),
            )
            data_import_succeeded = True
            dump_handoff(job_dir / "structured-data-handoff.json", handoff)
            s1 = handoff["s1"]
            result = service.record_data_understanding(
                project_id=project_id,
                expected_revision=result["revision"],
                datasource_inventory=s1["datasource_inventory"],
                schema_snapshot=s1["schema_snapshot"],
                data_profile=s1["data_profile"],
                relation_candidates=s1["relation_candidates"],
                evidence_sql=s1["evidence_sql"],
            )
            completed_workflow_stages.append("S1")
            structured_data = {
                "action": normalized_structured_action,
                "status": "READY_FOR_SEMANTIC_REVIEW",
                "database": "orion_source_data",
                "dataset_count": len(handoff["receipts"]),
                "sheet_count": sum(int(item["sheet_count"]) for item in handoff["receipts"]),
                "row_count": sum(int(item["row_count"]) for item in handoff["receipts"]),
                "handoff_artifact": "structured-data-handoff.json",
                "runtime_access_mode": "READ_ONLY",
                "next_stage": "S2",
                "next_action": "Harness 基于 S0 文档证据和 S1 数据画像生成业务语义候选",
            }
        except Exception as error:
            failed_stage = str((result or {}).get("current_stage") or "S1")
            structured_data = {
                "action": normalized_structured_action,
                "status": (
                    "WORKFLOW_GATE_FAILED" if data_import_succeeded else "IMPORT_FAILED"
                ),
                "data_import_status": "SUCCEEDED" if data_import_succeeded else "FAILED",
                "completed_workflow_stages": completed_workflow_stages,
                "failed_stage": failed_stage,
                "error": exception_detail(error),
                "retryable": data_import_succeeded,
            }
    status.update(
        {
            "status": "COMMITTED",
            "message": (
                "文件表格已全量导入并完成 S1 数据画像；工程停在 S2，等待 Harness 基于文档和数据库证据生成业务语义候选。"
                if structured_data["status"] == "READY_FOR_SEMANTIC_REVIEW"
                else (
                    (
                        "文件表格已全量导入，但后续工作流门禁未通过；"
                        f"已完成 {'/'.join(structured_data['completed_workflow_stages'])}，"
                        f"请修正 {structured_data['failed_stage']} 后从失败阶段续跑。"
                    )
                    if structured_data["status"] == "WORKFLOW_GATE_FAILED"
                    else (
                        "S0 已写入，但 PostgreSQL 结构化导入失败；"
                        "工程保留在当前安全阶段，请查明错误后重新发起导入批次。"
                    )
                    if structured_data["status"] == "IMPORT_FAILED"
                    else (
                        "已写入 S0，下一步进入 S1 数据理解。"
                        if intake_mode == "HYBRID"
                        else "已写入 S0，S1 已按范围留痕跳过，工程进入 S2。"
                    )
                )
            ),
            "committed_at": now(),
            "project_id": project_id,
            "workflow": result,
            "structured_data": structured_data,
            "reviewed_by": reviewed_by,
            "review_rationale": rationale.strip(),
        }
    )
    atomic_json(status_path, status)
    return status


def main() -> None:
    parser = argparse.ArgumentParser(description="ORION S0 文件资料批处理执行器")
    subparsers = parser.add_subparsers(dest="command", required=True)
    run_parser = subparsers.add_parser("run")
    run_parser.add_argument("--request-file", required=True, type=Path)
    commit_parser = subparsers.add_parser("commit")
    commit_parser.add_argument("--job-dir", required=True, type=Path)
    commit_parser.add_argument("--reviewed-by", required=True)
    commit_parser.add_argument("--rationale", required=True)
    commit_parser.add_argument("--accept-warnings", action="store_true")
    commit_parser.add_argument(
        "--structured-data-action",
        choices=("DOCUMENT_ONLY", "IMPORT"),
        default="DOCUMENT_ONLY",
    )
    rename_parser = subparsers.add_parser("rename-outputs")
    rename_parser.add_argument("--job-dir", required=True, type=Path)
    arguments = parser.parse_args()
    if arguments.command == "run":
        asyncio.run(run_job(arguments.request_file.resolve()))
        return
    if arguments.command == "rename-outputs":
        print(
            json.dumps(
                rename_job_markdown_outputs(arguments.job_dir),
                ensure_ascii=False,
            )
        )
        return
    result = commit_job(
        arguments.job_dir.resolve(),
        arguments.reviewed_by,
        arguments.rationale,
        arguments.accept_warnings,
        arguments.structured_data_action,
    )
    print(
        json.dumps(
            {
                "status": result["status"],
                "message": result["message"],
                "job_id": result["job_id"],
                "project_id": result["project_id"],
                "current_stage": result["workflow"]["current_stage"],
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
