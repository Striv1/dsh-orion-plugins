from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

from services.realtime_qa.query_capabilities import render_query_parameters
from services.realtime_qa.runtime_release import normalize_runtime_submission
from services.structured_data.pipeline import (
    StructuredDataImportError,
    StructuredDataPipeline,
    _business_sheet_label,
    _column_plans,
    _logical_source_reference,
)


def column(source_name: str, column_name: str, inferred_type: str = "text") -> dict:
    return {
        "column_index": 1,
        "source_name": source_name,
        "column_name": column_name,
        "inferred_type": inferred_type,
        "nullable": False,
        "non_null_count": 10,
        "distinct_count": 10,
        "minimum_value": "1",
        "maximum_value": "10",
    }


def test_column_inference_preserves_leading_zero_identifiers() -> None:
    columns = _column_plans(
        ["业务编号", "数量", "含税金额", "日期"],
        [
            ["00123", 1, "10.25", "2026-09-01"],
            ["00124", 2, "11.50", "2026-09-02"],
        ],
    )

    assert [item.inferred_type for item in columns] == [
        "text",
        "bigint",
        "numeric",
        "date",
    ]


def test_excel_bilingual_header_is_metadata_not_a_business_row(tmp_path: Path) -> None:
    source = tmp_path / "quality.xlsx"
    book = Workbook(write_only=True)
    sheet = book.create_sheet("lsj_dt_param")
    sheet.append(["batch_no", "actual_end_time", "param_value", "check_result"])
    sheet.append(["批次码", "出站时间", "参数值", "判定结果"])
    sheet.append(["B001", "2026-09-01 08:00:00", 3.5, "OK"])
    sheet.append(["B002", "2026-09-01 09:00:00", 4.5, "NG"])
    book.save(source)

    pipeline = StructuredDataPipeline("postgresql://unused", sample_rows=100)
    plans, rows = pipeline._excel_plans(source, "DS-TEST", "FILE-TEST")

    assert len(plans) == 1
    plan = plans[0]
    assert plan.header_row == 2
    assert plan.label_zh == "质检参数记录"
    assert [item.label_zh for item in plan.columns] == [
        "批次码",
        "出站时间",
        "参数值",
        "判定结果",
    ]
    assert [item.inferred_type for item in plan.columns] == [
        "text",
        "timestamp",
        "numeric",
        "text",
    ]
    imported = list(rows(plan))
    assert [item[0] for item in imported] == [3, 4]
    assert all("批次码" not in item[1] for item in imported)


def test_technical_sheet_names_receive_business_chinese_labels() -> None:
    assert _business_sheet_label(
        "lsj_dt_batch",
        [{"label_zh": "工序编码"}, {"label_zh": "设备编码"}],
    ) == "批次工序记录"
    assert _business_sheet_label(
        "lsj_dt_trace_batch",
        [{"label_zh": "电芯码"}],
    ) == "批次追溯记录"


@pytest.mark.parametrize("dimension", ["A1", "A1:B1048576", None])
def test_excel_streaming_ignores_bad_dimensions_without_full_memory_fallback(tmp_path, monkeypatch, dimension):
    import openpyxl

    original = tmp_path / "original.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["id", "amount"])
    sheet.append(["001", 2])
    sheet.append(["002", 3])
    book.save(original)
    source = tmp_path / "input.xlsx"
    with zipfile.ZipFile(original) as archive, zipfile.ZipFile(source, "w") as output:
        for item in archive.infolist():
            content = archive.read(item.filename)
            if dimension and item.filename == "xl/worksheets/sheet1.xml":
                content = re.sub(rb'<dimension ref="[^"]+"', f'<dimension ref="{dimension}"'.encode(), content)
            output.writestr(item, content)
    load = openpyxl.load_workbook
    modes = []

    def guarded_load(*args, **kwargs):
        modes.append(kwargs.get("read_only"))
        assert kwargs.get("read_only") is True, "大表不能回退到全内存读取"
        return load(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", guarded_load)
    pipeline = StructuredDataPipeline("postgresql://unused")
    plans, rows = pipeline._excel_plans(source, "DS-STREAM-TEST", "FILE-TEST")
    observed = list(rows(plans[0]))
    assert [row[0] for row in observed] == [2, 3]
    assert [row[1] for row in observed] == [("001", 2), ("002", 3)]
    assert modes == [True, True]


def test_upload_batch_id_is_not_part_of_logical_file_identity() -> None:
    first = _logical_source_reference(
        ".orion-s0-uploads/UPLOAD-AAAA/华东/经营数据.xlsx",
        "经营数据.xlsx",
    )
    second = _logical_source_reference(
        ".orion-s0-uploads/UPLOAD-BBBB/华东/经营数据.xlsx",
        "经营数据.xlsx",
    )

    assert first == second == "华东/经营数据.xlsx"


def test_word_and_pdf_tables_remain_document_evidence(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    pipeline = StructuredDataPipeline("postgresql://unused")
    imported: list[str] = []

    def import_document(**kwargs):
        imported.append(str(kwargs["document"]["source_type"]))
        return {"dataset_id": "DS-EXCEL"}

    monkeypatch.setattr(pipeline, "import_document", import_document)
    monkeypatch.setattr(
        pipeline,
        "build_handoff_for_datasets",
        lambda **kwargs: {"dataset_ids": kwargs["dataset_ids"]},
    )

    result = pipeline.import_documents(
        project_id="project-1",
        documents=[
            {"source_type": "DOCX", "structured_markdown": "| 字段 | 值 |"},
            {"source_type": "PDF", "structured_markdown": "| 字段 | 值 |"},
            {"source_type": "XLSX"},
        ],
        input_root=Path("/unused"),
    )

    assert imported == ["XLSX"]
    assert result == {"dataset_ids": ["DS-EXCEL"]}


def test_handoff_builds_s1_s2_s3_and_locatable_ontop_queries() -> None:
    pipeline = StructuredDataPipeline("postgresql://unused")
    receipt = {
        "dataset_id": "DS-1234567890ABCDEF",
        "project_id": "excel-project",
        "document_id": "FILE-001",
        "s0_document_id": "DOC-CONTENT-001",
        "source_name": "经营数据.xlsx",
        "source_type": "XLSX",
        "source_sha256": "sha256:" + "a" * 64,
        "source_uri": "minio://orion-workflow-artifacts/test.xlsx",
        "version": "a" * 16,
        "status": "READY",
        "sheet_count": 2,
        "row_count": 20,
        "imported_at": "2026-09-02T00:00:00+00:00",
        "promoted_at": "2026-09-02T00:00:00+00:00",
        "unchanged": False,
        "sheets": [
            {
                "sheet_id": "SH-CUSTOMERS",
                "sheet_index": 1,
                "source_name": "客户",
                "table_name": "ds_test_001",
                "view_name": "current_file_001_schema",
                "row_count": 10,
                "column_count": 2,
                "columns": [
                    column("customer_id", "customer_id"),
                    column("参数值", "param_value", "numeric"),
                ],
            },
            {
                "sheet_id": "SH-ORDERS",
                "sheet_index": 2,
                "source_name": "订单",
                "table_name": "ds_test_002",
                "view_name": "current_file_002_schema",
                "row_count": 10,
                "column_count": 1,
                "columns": [column("customer_id", "customer_id")],
            },
        ],
    }

    handoff = pipeline.build_handoff(project_id="excel-project", receipts=[receipt])

    assert handoff["s1"]["data_profile"]["total_rows"] == 20
    assert len(handoff["s1"]["relation_candidates"]) == 1
    relation_evidence_id = handoff["s1"]["relation_candidates"][0]["source_refs"][-1]
    relation_evidence = next(
        item
        for item in handoff["s1"]["evidence_sql"]
        if item["id"] == relation_evidence_id
    )
    assert relation_evidence["status"] == "NEEDS_EXECUTION"
    assert "executed_via" not in relation_evidence
    assert any(
        item["kind"] == "OBJECT_PROPERTY"
        and item["status"] == "NEEDS_HUMAN_CONFIRMATION"
        for item in handoff["s2"]["ontology_candidates"]
    )
    runtime = normalize_runtime_submission(
        handoff["s3"]["realtime_runtime"],
        intake_mode="HYBRID",
        require_explicit_capabilities=True,
    )
    assert runtime is not None
    query = next(iter(runtime["ontop_queries"].values()))
    assert "?dataset_id" in query
    assert "?source_row" in query
    assert ":orionDatasetId" in runtime["mapping_obda"]
    assert runtime["database_access_mode"] == "READ_ONLY"
    assert runtime["document_query_capabilities"] == [
        "current_full_text_search",
        "reviewed_entity_evidence",
    ]
    count_name = next(
        name
        for name, capability in runtime["query_capabilities"].items()
        if name.startswith("count_")
        and any(
            spec["description_zh"] == "参数值的筛选值"
            for spec in capability["parameters"].values()
        )
    )
    capability = runtime["query_capabilities"][count_name]
    value_parameter = next(
        name
        for name, spec in capability["parameters"].items()
        if spec["description_zh"] == "参数值的筛选值"
    )
    operator_parameter = next(
        name
        for name, spec in capability["parameters"].items()
        if spec["description_zh"] == "参数值的比较方式"
    )
    rendered = render_query_parameters(
        runtime["ontop_queries"][count_name],
        {value_parameter: 1, operator_parameter: "GT"},
        capability,
    )
    assert "COUNT(DISTINCT ?entity)" in rendered
    assert "> 1" in rendered
    assert "{{" not in rendered
    assert any(name.startswith("group_") for name in runtime["query_capabilities"])
    assert any(name.startswith("stats_") for name in runtime["query_capabilities"])
    assert any(name.startswith("relation_") for name in runtime["query_capabilities"])
    assert all(
        capability["result_fields"] and capability["question_examples"]
        for capability in runtime["query_capabilities"].values()
    )
    assert all(
        any("\u3400" <= character <= "\u9fff" for character in item["target_label_zh"])
        for item in handoff["s3"]["mapping_draft"]["mappings"]
    )
    class_names = [
        item["name"]
        for item in handoff["s2"]["ontology_candidates"]
        if item["kind"] == "CLASS"
    ]
    assert len(class_names) == len(set(class_names)) == 2


def test_import_rejects_a_source_that_changed_after_s0(tmp_path: Path) -> None:
    source = tmp_path / "changed.csv"
    source.write_text("id,name\n1,changed\n", encoding="utf-8")
    pipeline = StructuredDataPipeline("postgresql://unused")

    with pytest.raises(StructuredDataImportError, match="SHA-256 已变化"):
        pipeline.import_document(
            project_id="project-1",
            document={
                "document_id": "DOC-OLD",
                "source_name": source.name,
                "source_type": "CSV",
                "source_sha256": "sha256:" + "0" * 64,
                "source_path": source.name,
            },
            input_root=tmp_path,
        )
